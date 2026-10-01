# Part of Odoo. See LICENSE file for full copyright and licensing details.

from datetime import date
from io import BytesIO
from openpyxl import load_workbook

from odoo import Command
from odoo.tests import tagged

from .common import TestPayrollCommon


@tagged('post_install_l10n', 'post_install', '-at_install')
class TestPaymentReport(TestPayrollCommon):

    def test_excel_sheet_format_eta_form2(self):
        """Check ETA Form2 report format"""

        self.employee.write({
            'identification_id': '212121211',
            'passport_id': '342324fsad',
            'l10n_eg_ssn': 'fdjsakfjas',
            'job_id': self.env['hr.job'].create({
                'name': 'Developer',
            }),
            'resource_calendar_id': self.env.ref('l10n_eg_hr_payroll.resource_calendar_def_40h').id,
        })
        self.employee_payslip = self.env['hr.payslip'].create({
            'employee_id': self.employee.id,
            'version_id': self.employee.version_id.id,
            'struct_id': self.employee.structure_id.id,
            'date_from': date(2025, 5, 1),
            'date_to': date(2025, 5, 31),
        })
        self.employee_payslip.compute_sheet()
        self.employee_payslip.action_payslip_done()

        self.env['hr.payroll.payment.report.wizard'].create({
            'export_format': 'l10n_eg_eta_form2',
            'payslip_ids': [Command.link(self.employee_payslip.id)],
        }).generate_payment_report()

        workbook = load_workbook(BytesIO(self.employee_payslip.payment_report))
        sheet = workbook.active
        headers_x_values_map = {
            'Serial Number': 1,
            'Employee Name': self.employee.name,
            'National Number/ID': self.employee.identification_id,
            'Passport Number': self.employee.passport_id,
            'Insurance State': 1,
            'Insurance Number': self.employee.l10n_eg_ssn,
            'Job Position': self.employee.job_title,
            'Total Amount Paid to Employee': self.employee_payslip.net_wage,
        }
        for col, (header, value) in enumerate(headers_x_values_map.items()):
            self.assertEqual(sheet.cell(1, col + 1).value, header)
            self.assertEqual(sheet.cell(2, col + 1).value, value)
