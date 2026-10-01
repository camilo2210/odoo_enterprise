# Part of Odoo. See LICENSE file for full copyright and licensing details.

from io import BytesIO
from openpyxl import load_workbook

from odoo.tests import common, tagged

from .common import TestPayrollCommon


@tagged('post_install_l10n', 'post_install', '-at_install')
class TestEgyptianHrEmployee(TestPayrollCommon, common.HttpCase):

    def test_errors_on_nosi_form2(self):
        # company doesn't have required fields
        nosi_form2 = self.env['l10n.eg.nosi.form2.wizard'].create({
            'employee_ids': [(6, 0, self.employee.ids)],
        })
        with self.assertRaises(Exception):
            nosi_form2.action_generate_report()
        self.company.write({
            'vat': '587493026',
            'l10n_eg_nosi_code': 'NOSI123456',
        })

        # employee doesn't have required fields
        with self.assertRaises(Exception):
            nosi_form2.action_generate_report()
        self.employee.write({
            'l10n_eg_ssn': 'l10n_eg_ssn12345',
            'identification_id': 'identification12345',
            'l10n_eg_social_insurance_reference': '400',
        })
        nosi_form2.action_generate_report()

    def test_excel_sheet_format_nosi_form2(self):
        # Open the nosi form 2 and check whether the headers and employee values are present or not.
        self.company.write({
            'vat': '587493026',
            'l10n_eg_nosi_code': 'NOSI123456',
        })
        self.employee.write({
            'l10n_eg_ssn': 'l10n_eg_ssn123456',
            'identification_id': 'identification_id123456',
            'l10n_eg_social_insurance_reference': '500',
        })
        self.authenticate('admin', 'admin')

        nosi_form2 = self.env['l10n.eg.nosi.form2.wizard'].create({
            'employee_ids': [(6, 0, self.employee.ids)],
        })
        returned_response = nosi_form2.action_generate_report()
        response = self.url_open(returned_response.get('url'))

        # check for response successful and recieved spreadsheet file
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            response.headers.get('Content-Type', ''),
        )
        self.assertIn(
            'Nosi_form_2.xlsx',
            response.headers.get('Content-Disposition', ''),
        )

        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active

        # check for headers in spreadsheet
        self.assertEqual(sheet.cell(1, 1).value, 'Company Details')
        for col, cell_header in enumerate(['NOSI Registration', 'Company Name', 'Company VAT ID', 'Company Address']):
            self.assertEqual(sheet.cell(2, col + 1).value, cell_header)
        self.assertEqual(sheet.cell(5, 1).value, 'Employees Details')
        for col, cell_header in enumerate([
            'Insurance Number',
            'Name',
            'ID/National Number',
            'Joining Date',
            'NOSI Contribution Amount',
            'Total Salary',
        ]):
            self.assertEqual(sheet.cell(6, col + 1).value, cell_header)

        # check for values in spreadsheet
        self.assertEqual(sheet.cell(3, 2).value, self.company.name)
        employee_row = [
            self.employee.l10n_eg_ssn, self.employee.name, self.employee.identification_id, '01/01/2016',
            str(int(self.employee.l10n_eg_social_insurance_reference)), str(int(self.employee.current_version_id._get_contract_wage())),
        ]
        for col, data in enumerate(employee_row):
            self.assertEqual(str(sheet.cell(7, col + 1).value), data)
