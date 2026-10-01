from datetime import date, datetime
import csv
import io
from openpyxl import load_workbook
from odoo.fields import Command
from odoo.tests import tagged
from .common import TestPayrollCommon


@tagged('post_install_l10n', 'post_install', '-at_install', 'payroll_meal_vouchers')
class TestPayrollMealVouchers(TestPayrollCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.work_entry_type = cls.env['hr.work.entry.type'].create({
            'name': 'Legal Leaves',
            'code': 'Legal Leaves',
            'count_as': 'absence',
            'requires_allocation': False,
            'unit_of_measure': 'day',
        })

    def _get_meal_voucher_quantity(self, payslip):
        return payslip.line_ids.filtered(lambda line: line.salary_rule_id.code == 'MEAL_V_EMP').quantity or 0

    def _create_report_for_export(self):
        employee = self.employee_john
        employee.niss = '03062411377'
        employee.meal_voucher_amount = 7.5
        self.create_and_validate_payslips(employee, 2023, [3])
        report = self.env['l10n_be.meal.vouchers.report'].create({
            'company_id': self.belgian_company.id,
            'month': '3',
            'year': '2023',
        })
        self.env.flush_all()
        self.env.invalidate_all()
        line = report.meal_vouchers_line_ids.filtered(lambda l: l.employee_id == employee)
        self.assertEqual(len(line), 1)
        return report, employee, line

    def _read_xlsx_rows(self, xlsx_content):
        workbook = load_workbook(filename=io.BytesIO(xlsx_content), read_only=True, data_only=True)
        try:
            worksheet = workbook['Meal Vouchers'] if 'Meal Vouchers' in workbook.sheetnames else workbook.worksheets[0]
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return rows

    def test_meal_vouchers_computation_with_postponed(self):
        # -In case of postponed Meal vouchers, it may happen that the value of (patronal part + employee part) change.
        # -Example:
        # Feb → MV amount 8€
        # March → MV amount 10€
        # -Employee received too many MV in Feb because some absences where encoded after the order of the MV.
        # -In that case, it will be postponed on March but the employee will be charged for a MV of a value of 10€ while he received meal voucher of 8€.
        # The total has been adjusted by considering the original value for meal_voucher_amount in case of postponed MV.
        employee = self.employee_john
        employee.niss = '03062411377'
        employee.write({'meal_voucher_amount': 8})
        feb_payslip = self.env['hr.payslip'].create({
            'employee_id': employee.id,
            'version_id': employee.version_id.id,
            'date_from': date(2023, 2, 1),
            'date_to': date(2023, 2, 28),
            'company_id': self.belgian_company.id,
        })
        feb_payslip.compute_sheet()
        feb_payslip.action_payslip_done()
        feb_payslip.action_payslip_paid()

        feb_report = self.env['l10n_be.meal.vouchers.report'].create({
            'company_id': self.belgian_company.id,
            'month': '2',
            'year': '2023',
        })
        feb_report.action_set_state('done')
        self.env.flush_all()
        self.env.invalidate_all()
        feb_meal_vouchers_line = feb_report.meal_vouchers_line_ids
        self.assertEqual(len(feb_meal_vouchers_line), 1)
        self.assertEqual(feb_meal_vouchers_line.entitlement, 12)
        self.assertEqual(feb_meal_vouchers_line.postponed, 0)
        self.assertEqual(feb_meal_vouchers_line.total, 12)
        self.assertEqual(feb_meal_vouchers_line.total_value, employee.meal_voucher_amount * feb_meal_vouchers_line.total)

        leave = self.env['hr.leave'].create({
            'employee_id': employee.id,
            'work_entry_type_id': self.work_entry_type.id,
            'request_date_from': datetime(2023, 2, 27),
            'request_date_to': datetime(2023, 2, 28),
        })
        leave.action_approve()

        feb_payslip.refund_sheet()
        refund_payslip = self.env['hr.payslip'].search([
            ('origin_payslip_id', '=', feb_payslip.id),
            ('credit_note', '=', True),
        ])
        self.assertEqual(len(refund_payslip), 1)
        leave.payslip_state = 'normal'

        new_feb_payslip = self.env['hr.payslip'].create({
            'employee_id': employee.id,
            'version_id': employee.version_id.id,
            'date_from': date(2023, 2, 1),
            'date_to': date(2023, 2, 28),
            'company_id': self.belgian_company.id,
            'origin_payslip_id': feb_payslip.id,
        })
        new_feb_payslip.compute_sheet()
        new_feb_payslip.action_payslip_done()
        new_feb_payslip.action_payslip_paid()
        leave_line = new_feb_payslip.worked_days_line_ids.filtered(lambda line: line.code == 'Legal Leaves')
        edited_meal_line = new_feb_payslip.line_ids.filtered(lambda line: line.code == 'MEAL_V_EMP')
        self.assertEqual(leave_line.number_of_days, 2)
        self.assertEqual(len(edited_meal_line), 1)
        self.assertEqual(edited_meal_line.quantity, 10)

        feb_version = employee.version_id
        feb_version.write({
            'contract_date_end': date(2023, 2, 28)
        })
        march_version = feb_version.copy({
            'date_version': date(2023, 3, 1),
            'contract_date_start': date(2023, 3, 1),
            'contract_date_end': False,
            'meal_voucher_amount': 10,
        })
        march_payslip = self.env['hr.payslip'].create({
            'employee_id': employee.id,
            'version_id': march_version.id,
            'date_from': date(2023, 3, 1),
            'date_to': date(2023, 3, 31),
            'company_id': self.belgian_company.id,
        })
        march_payslip.compute_sheet()
        march_payslip.action_payslip_done()
        march_payslip.action_payslip_paid()

        march_report = self.env['l10n_be.meal.vouchers.report'].create({
            'company_id': self.belgian_company.id,
            'month': '3',
            'year': '2023',
        })
        self.env.flush_all()
        self.env.invalidate_all()
        march_line = march_report.meal_vouchers_line_ids.filtered(lambda l: l.employee_id == employee)
        self.assertEqual(len(march_line), 1)

        march_meal_vouchers = self._get_meal_voucher_quantity(march_payslip)
        old_feb_meal_vouchers = self._get_meal_voucher_quantity(feb_payslip)
        new_feb_meal_vouchers = self._get_meal_voucher_quantity(new_feb_payslip)

        self.assertEqual(march_line.entitlement, march_meal_vouchers)
        self.assertEqual(march_line.postponed, old_feb_meal_vouchers - new_feb_meal_vouchers)
        self.assertEqual(march_line.total, march_meal_vouchers - old_feb_meal_vouchers + new_feb_meal_vouchers)
        self.assertAlmostEqual(march_line.total_value, (employee.meal_voucher_amount * march_meal_vouchers) - (feb_meal_vouchers_line.value * march_line.postponed), places=2)
        self.assertAlmostEqual(march_line.value, march_line.total_value / march_line.total, places=2)

    def test_meal_vouchers_postponed_across_branch_companies(self):
        # Meal vouchers postponed by a payslip correction must be forwarded to a meal
        # voucher report even when the report's company is only an ancestor (branch root)
        # of the payslip's company
        root_company, branch_company = self.multibranch_company[0], self.multibranch_company[1]
        self.env.user.company_ids |= (root_company | branch_company)
        self.env = self.env(context=dict(
            self.env.context,
            allowed_company_ids=self.env.context.get('allowed_company_ids', []) + (root_company | branch_company).ids,
        ))
        employee = self.create_employee({
            'name': 'Employee Branch',
            'company_id': branch_company.id,
            'contract_date_start': date(2026, 1, 1),
            'meal_voucher_amount': 8,
            'meal_voucher_employee_share': 1.09,
        })

        jan_payslip = self.env['hr.payslip'].create({
            'name': f'Meal vouchers - {employee.name}',
            'employee_id': employee.id,
            'version_id': employee.version_id.id,
            'date_from': date(2026, 1, 1),
            'date_to': date(2026, 1, 31),
            'company_id': branch_company.id,
        })
        jan_payslip.compute_sheet()
        jan_payslip.action_payslip_done()
        jan_payslip.action_payslip_paid()
        original_quantity = self._get_meal_voucher_quantity(jan_payslip)

        jan_report = self.env['l10n_be.meal.vouchers.report'].with_company(root_company).create({
            'company_id': root_company.id,
            'month': '1',
            'year': '2026',
        })
        self.env.flush_all()
        self.env.invalidate_all()
        jan_line = jan_report.meal_vouchers_line_ids.filtered(lambda l: l.employee_id == employee)
        self.assertEqual(len(jan_line), 1, "The branch company employee should appear in the root company's report.")
        # Close the January report, as it would be in a normal workflow, before the payslip
        # gets corrected. Once linked to a "done" report, the original payslip no longer leaks
        # into every future (not yet done) report as a phantom postponed entitlement.
        jan_report.action_set_state('done')

        leave = self.env['hr.leave'].create({
            'employee_id': employee.id,
            'work_entry_type_id': self.work_entry_type.id,
            'request_date_from': datetime(2026, 1, 26),
            'request_date_to': datetime(2026, 1, 28),
        })
        leave.action_approve()

        jan_payslip.refund_sheet()
        leave.payslip_state = 'normal'

        corrected_jan_payslip = self.env['hr.payslip'].create({
            'name': f'Meal vouchers correction - {employee.name}',
            'employee_id': employee.id,
            'version_id': employee.version_id.id,
            'date_from': date(2026, 1, 1),
            'date_to': date(2026, 1, 31),
            'company_id': branch_company.id,
            'origin_payslip_id': jan_payslip.id,
        })
        corrected_jan_payslip.compute_sheet()
        corrected_jan_payslip.action_payslip_done()
        corrected_jan_payslip.action_payslip_paid()
        corrected_quantity = self._get_meal_voucher_quantity(corrected_jan_payslip)
        self.assertLess(corrected_quantity, original_quantity, "The leave should reduce the meal vouchers entitlement.")

        feb_payslip = self.env['hr.payslip'].create({
            'name': f'Meal vouchers - {employee.name}',
            'employee_id': employee.id,
            'version_id': employee.version_id.id,
            'date_from': date(2026, 2, 1),
            'date_to': date(2026, 2, 28),
            'company_id': branch_company.id,
        })
        feb_payslip.compute_sheet()
        feb_payslip.action_payslip_done()
        feb_payslip.action_payslip_paid()

        feb_report = self.env['l10n_be.meal.vouchers.report'].with_company(root_company).create({
            'company_id': root_company.id,
            'month': '2',
            'year': '2026',
        })
        self.assertEqual(feb_report.state, 'draft', "The bug only reproduces when the report is not marked as done.")
        self.env.flush_all()
        self.env.invalidate_all()
        feb_line = feb_report.meal_vouchers_line_ids.filtered(lambda l: l.employee_id == employee)
        self.assertEqual(len(feb_line), 1)
        self.assertEqual(
            feb_line.postponed, original_quantity - corrected_quantity,
            "Meal vouchers postponed by a branch company's payslip correction should be forwarded "
            "to the parent company's report, even when that report is not marked as done.")

    def test_meal_vouchers_computation_per_hours(self):
        employee = self.create_employee_with_benefits({
            'name': 'Meal Vouchers per Hours',
            'resource_calendar_id': self.resource_calendar_40.id,
            'reference_calendar_id': self.belgian_company.resource_calendar_id.id,
            'date_version': date(2026, 2, 1),
            'contract_date_start': date(2026, 2, 1),
        })
        payslip = self.env['hr.payslip'].create({
            'employee_id': employee.id,
            'version_id': employee.version_id.id,
            'date_from': date(2026, 2, 1),
            'date_to': date(2026, 2, 28),
            'company_id': self.belgian_company.id,
        })

        self.assertEqual(employee.meal_voucher_calculation_method, 'days')
        payslip.compute_sheet()
        meal_voucher_line = payslip.line_ids.filtered(lambda line: line.code == 'MEAL_V_EMP')
        # February 2026 has 20 eligible working days: 20 vouchers x EUR 1.09 employee share.
        self.assertRecordValues(meal_voucher_line, [{
            'amount': 1.09,
            'quantity': 20,
            'rate': -100,
            'total': -21.80,
        }])

        employee.meal_voucher_calculation_method = 'hours'
        payslip.compute_sheet()
        meal_voucher_line = payslip.line_ids.filtered(lambda line: line.code == 'MEAL_V_EMP')
        # 160 eligible hours / 7.6 reference hours per day = 21.05, rounded to 21 vouchers.
        self.assertRecordValues(meal_voucher_line, [{
            'amount': 1.09,
            'quantity': 21,
            'rate': -100,
            'total': -22.89,
        }])

    def test_meal_vouchers_export(self):
        report, employee, line = self._create_report_for_export()

        report.with_context().action_generate_meal_vouchers_report()

        self.assertEqual(report.state, 'ready')
        self.assertEqual(report.export_filename_csv, '2023_03_MealVoucher.csv')
        self.assertTrue(report.export_file_csv)
        self.assertEqual(report.export_filename_xlsx, '2023_03_MealVoucher.xlsx')
        self.assertTrue(report.export_file_xlsx)

        csv_content = report.export_file_csv.content.decode('utf-8')
        rows = list(csv.reader(io.StringIO(csv_content)))

        self.assertEqual(rows[0], ['Name', 'NISS', 'Value', 'Quantity', 'Total'])
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], employee.name)
        self.assertEqual(rows[1][1], employee.niss)
        self.assertAlmostEqual(float(rows[1][2]), employee.meal_voucher_amount, places=2)
        self.assertAlmostEqual(float(rows[1][3]), line.total, places=2)
        self.assertAlmostEqual(float(rows[1][4]), line.total_value, places=2)

        xlsx_content = report.export_file_xlsx.content
        rows = self._read_xlsx_rows(xlsx_content)

        self.assertEqual(rows[0], ['Name', 'NISS', 'Value', 'Quantity', 'Total'])
        self.assertEqual(rows[1][0], employee.name)
        self.assertEqual(rows[1][1], employee.niss)
        self.assertAlmostEqual(float(rows[1][2]), employee.meal_voucher_amount, places=2)
        self.assertAlmostEqual(float(rows[1][3]), line.total, places=2)
        self.assertAlmostEqual(float(rows[1][4]), line.total_value, places=2)

    def test_employee_departure_with_legal_leaves_meal_vouchers(self):
        """
        Checks that the right number of meal vouchers is given to an employee when they leave the company.
        Should be equal to number_of_days_per_week * 52 - 10 (public holidays) - time-off allocation (paid time off + extra legal)
        """
        allocation_PTO, allocation_extra = self.env['hr.leave.allocation'].create([{
            'name': 'Test Allocation PTO',
            'date_from': date(2025, 1, 20),
            'employee_id': self.employee_a.id,
            'work_entry_type_id': self.env.ref('hr_work_entry.be_work_entry_type_legal_leave').id,
            'number_of_days': 4,
        }, {
            'name': 'Test Allocation extra',
            'date_from': date(2025, 1, 20),
            'employee_id': self.employee_a.id,
            'work_entry_type_id': self.env.ref('hr_work_entry.l10n_be_work_entry_type_extra_legal').id,
            'number_of_days': 8,
        }])
        allocation_PTO.action_approve()
        allocation_extra.action_approve()
        departure_notice = self.env['hr.employee.departure'].create({
            'employee_id': self.employee_a.id,
            'departure_reason_id': self.env.ref('hr.departure_fired').id,
            'dismissal_date': date(2025, 12, 20),
            'l10n_be_notice_respect': 'without',
            'departure_description': "Oh no you're fired",
        })
        departure_notice.action_register()
        departure_payslips = departure_notice._generate_termination_payslip()
        departure_payslips = self.env['hr.payslip'].search([('id', 'in', departure_payslips.ids)])
        struct_id = self.env.ref('l10n_be_hr_payroll.hr_payroll_structure_cp200_employee_termination_fees')
        termination_fees = departure_payslips.filtered(lambda dep: dep.struct_id == struct_id)
        termination_fees.version_id.meal_voucher_amount = 8
        termination_fees.compute_sheet()
        meal_vouchers_line = termination_fees.line_ids.filtered(lambda l: l.name == 'Meal Voucher')
        days_per_week = termination_fees.version_id.resource_calendar_id.days_per_week
        days_off = sum((allocation_PTO | allocation_extra).mapped('number_of_days'))
        self.assertEqual(meal_vouchers_line.quantity, (days_per_week * 52) - 10 - days_off)
        self.assertEqual(meal_vouchers_line.amount, termination_fees.version_id.meal_voucher_amount - termination_fees.version_id.meal_voucher_employee_share)

    def test_part_time_employee_departure_with_legal_leaves_meal_vouchers(self):
        """
        Checks that the right number of meal vouchers is given to an employee when they leave the company.
        Should be equal to number_of_days_per_week * 52 - 10 (public holidays) - time-off allocation (paid time off + extra legal)
        """
        self.employee_georges.resource_calendar_id.write({'days_per_week': 4})
        allocation_PTO, allocation_extra = self.env['hr.leave.allocation'].create([{
            'name': 'Test Allocation PTO',
            'date_from': date(2025, 1, 20),
            'employee_id': self.employee_georges.id,
            'work_entry_type_id': self.env.ref('hr_work_entry.be_work_entry_type_legal_leave').id,
            'number_of_days': 4,
        }, {
            'name': 'Test Allocation extra',
            'date_from': date(2025, 1, 20),
            'employee_id': self.employee_georges.id,
            'work_entry_type_id': self.env.ref('hr_work_entry.l10n_be_work_entry_type_extra_legal').id,
            'number_of_days': 20,
        }])
        allocation_PTO.action_approve()
        allocation_extra.action_approve()
        departure_notice = self.env['hr.employee.departure'].create({
            'employee_id': self.employee_georges.id,
            'departure_reason_id': self.env.ref('hr.departure_fired').id,
            'dismissal_date': date(2025, 12, 20),
            'l10n_be_notice_respect': 'without',
            'departure_description': "Oh no you're fired",
        })
        departure_notice.action_register()
        departure_payslips = departure_notice._generate_termination_payslip()
        departure_payslips = self.env['hr.payslip'].search([('id', 'in', departure_payslips.ids)])
        struct_id = self.env.ref('l10n_be_hr_payroll.hr_payroll_structure_cp200_employee_termination_fees')
        termination_fees = departure_payslips.filtered(lambda dep: dep.struct_id == struct_id)
        termination_fees.version_id.meal_voucher_amount = 8
        termination_fees.compute_sheet()
        meal_vouchers_line = termination_fees.line_ids.filtered(lambda l: l.name == 'Meal Voucher')
        days_per_week = termination_fees.version_id.resource_calendar_id.days_per_week
        days_off = sum((allocation_PTO | allocation_extra).mapped('number_of_days'))
        self.assertEqual(meal_vouchers_line.quantity, (days_per_week * 52) - 10 - days_off)
        self.assertEqual(meal_vouchers_line.amount, termination_fees.version_id.meal_voucher_amount - termination_fees.version_id.meal_voucher_employee_share)

    def test_include_child_branches_in_report(self):
        companies = self.belgian_company | self.multibranch_company
        # company 0 -> root / no child
        # company 1 (is parent of) company 2 (is parent of) company 3
        target_company = companies[1]

        employees = self.create_employee([
            {
                'name': f'Employee {company.name}',
                'company_id': company.id,
                'contract_date_start': date(2026, 1, 1),
                'meal_voucher_amount': 10.,
                'meal_voucher_employee_share': 1.09,
            } for company in companies
        ])

        self.create_and_validate_payslips(employees=employees, year=2026, months=[1])

        report = self.env['l10n_be.meal.vouchers.report'].with_company(target_company).create({
            'year': 2026,
            'month': 1,
        })

        expected_companies = companies[1:]
        self.assertEqual(report.branch_ids, expected_companies,
                         "Report should include all three levels of the hierarchy.")

        expected_reported_employees = employees[1:]
        reported_employees = report.meal_vouchers_line_ids.mapped('employee_id')
        self.assertEqual(reported_employees, expected_reported_employees,
                         "Report should include all three employees from companies with indexes 1-2-3")

        self.assertEqual(len(report.meal_vouchers_line_ids), 3,
                         "There should be exactly one line per employee across all branches.")

    def _create_union_mission_payslip(self, category_options):
        """ Take 2 union mission days (021.00) in February 2023 and return the payslip of that month. """
        employee = self.employee_john
        union_mission_type = self.env.ref('hr_work_entry.l10n_be_work_entry_type_union_mission')
        union_mission_type.requires_allocation = False
        leave = self.env['hr.leave'].create({
            'employee_id': employee.id,
            'work_entry_type_id': union_mission_type.id,
            'request_date_from': datetime(2023, 2, 27),
            'request_date_to': datetime(2023, 2, 28),
            'category_options_ids': [Command.set(category_options.ids)],
        })
        leave.action_approve()

        payslip = self.env['hr.payslip'].create({
            'employee_id': employee.id,
            'version_id': employee.version_id.id,
            'date_from': date(2023, 2, 1),
            'date_to': date(2023, 2, 28),
            'company_id': self.belgian_company.id,
        })
        payslip.compute_sheet()
        return payslip

    def test_meal_vouchers_union_mission_without_option(self):
        # Union mission days don't grant a meal voucher by default: 12 - 2 days
        payslip = self._create_union_mission_payslip(self.env['hr.salary.rule.category'])
        self.assertEqual(self._get_meal_voucher_quantity(payslip), 10)

    def test_meal_vouchers_union_mission_with_option(self):
        # The meal voucher is an optional category of 021.00: when the HR user ticks it on the time off, those days grant a meal voucher as well
        meal_voucher_category = self.env.ref('l10n_be_hr_payroll.MEAL_VOUCHER')
        self.assertIn(
            meal_voucher_category,
            self.env.ref('hr_work_entry.l10n_be_work_entry_type_union_mission').optional_category_ids)
        payslip = self._create_union_mission_payslip(meal_voucher_category)
        self.assertEqual(self._get_meal_voucher_quantity(payslip), 12)

        payslip.employee_id.meal_voucher_calculation_method = 'hours'
        payslip.compute_sheet()
        # 76 hours (including the 2 union mission days) / 7.6 hours per day = 10
        self.assertEqual(self._get_meal_voucher_quantity(payslip), 10)
