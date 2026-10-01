# Part of Odoo. See LICENSE file for full copyright and licensing details.

import re
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from odoo.addons.hr_payroll_account.tests.common import TestPayslipValidationCommon
from odoo.tests.common import tagged
from odoo.exceptions import UserError


@tagged('post_install_l10n', 'post_install', '-at_install', 'nosi')
class TestNosiExport(TestPayslipValidationCommon):
    """
    Test suite for the NOSI Form 1 and Form 6 XLSX export functionality using the wizard.
    """
    _test_user_groups = None  # FIXME list needed groups

    @classmethod
    @TestPayslipValidationCommon.setup_country('eg')
    def setUpClass(cls):
        super().setUpClass()
        cls._setup_common(
            country=cls.env.ref('base.eg'),
            structure=cls.env.ref('l10n_eg_hr_payroll.hr_payroll_structure_eg_employee_salary'),
            structure_type=cls.env.ref('l10n_eg_hr_payroll.structure_type_employee_eg'),
        )

        cls.country_eg = cls.env.ref('base.eg')
        cls.address = cls.env['res.partner'].create({
            'name': 'Test Address Partner',
            'street': '123 Main St',
            'city': 'Cairo',
            'country_id': cls.country_eg.id,
        })
        cls.departure_reason = cls.env['hr.departure.reason'].create({'name': 'Resignation'})
        cls.software_developer_job = cls.env['hr.job'].sudo().create({'name': 'SDE'})
        cls.manager_job = cls.env['hr.job'].sudo().create({'name': 'Manager'})
        cls.company.l10n_eg_nosi_code = 'NOSI-12345'

        # --- Create Employees with all necessary data ---
        cls.employee_1 = cls.env['hr.employee'].create({
            'name': 'Mohamed Ali',
            'company_id': cls.company.id,
            'l10n_eg_ssn': 'SSN123',
            'identification_id': 'ID123',
            'country_id': cls.country_eg.id,
            'job_id': cls.software_developer_job.id,
            'job_title': 'Developer',
            'l10n_eg_nosi_registration_start_date': date(2023, 1, 1),
            'wage': 5000.0,
            'l10n_eg_social_insurance_reference': 4500.0,
            'private_phone': '111-222-333',
            'private_email': 'mohamed.ali@example.com',
            'address_id': cls.address.id,
            'contract_date_start': date(2020, 1, 1),
            'l10n_eg_housing_allowance': 100,
            'l10n_eg_transportation_allowance': 110,
            'l10n_eg_other_allowances': 80,
        })
        cls.employee_2 = cls.env['hr.employee'].create({
            'name': 'Ahmed Ali',
            'company_id': cls.company.id,
            'l10n_eg_ssn': 'SSN456',
            'identification_id': 'ID456',
            'country_id': cls.country_eg.id,
            'job_id': cls.manager_job.id,
            'job_title': 'Manager',
            'l10n_eg_nosi_registration_start_date': date(2023, 2, 1),
            'wage': 8000.0,
            'l10n_eg_social_insurance_reference': 7000.0,
            'private_phone': '444-555-666',
            'private_email': 'ahmed.ali@example.com',
            'address_id': cls.address.id,
            'contract_date_start': date(2020, 1, 1),
        })

        cls.employee_1.structure_id = cls.env.ref('l10n_eg_hr_payroll.hr_payroll_structure_eg_employee_salary')
        cls.employee_2.structure_id = cls.env.ref('l10n_eg_hr_payroll.hr_payroll_structure_eg_employee_salary')

    def _create_wizard(self, form_type, employee_ids):
        """Helper to create a NOSI export wizard with given parameters."""
        return self.env[f'l10n.eg.nosi.{form_type}.wizard'].create({
            'employee_ids': [(6, 0, employee_ids)],
            'company_id': self.company.id,
        })

    def _get_attachment_from_action(self, action):
        """Helper to extract and validate the attachment from a download action."""
        self.assertIsInstance(action, dict)
        self.assertEqual(action.get('type'), 'ir.actions.act_url')
        self.assertIn('/web/content/', action.get('url', ''))

        # Extract attachment ID from the URL
        match = re.search(r'/web/content/(\d+)', action['url'])
        self.assertTrue(match, "Could not find attachment ID in the URL")
        attachment_id = int(match.group(1))

        attachment = self.env['ir.attachment'].browse(attachment_id)
        self.assertTrue(attachment.exists(), "Attachment was not created in the database.")
        self.assertTrue(attachment.raw, "Attachment was created but has no data.")
        return attachment

    def _read_xlsx_content(self, attachment):
        """Helper to read XLSX file content from attachment."""
        workbook = load_workbook(BytesIO(attachment.raw))
        return workbook.active

    def _validate_form_headers(self, sheet, expected_headers):
        """Validate Form headers are correct."""
        # Headers are in row 1 (Excel is 1-indexed)
        for col_idx, expected_header in enumerate(expected_headers, start=1):
            actual_header = sheet.cell(row=1, column=col_idx).value
            self.assertEqual(actual_header, expected_header, f"Header mismatch at column {col_idx}: expected '{expected_header}', got '{actual_header}'")

    def _validate_employee_row(self, sheet, row_idx, employee, expected_results):
        """Validate a single employee row in the sheet."""
        row_data = [sheet.cell(row=row_idx, column=col).value for col in range(1, len(expected_results) + 1)]

        for col_idx, (actual, expected) in enumerate(zip(row_data, expected_results), start=1):
            if isinstance(expected, (int, float)):
                self.assertAlmostEqual(float(actual) if actual else 0, float(expected), places=2,
                                     msg=f"Row {row_idx}, Col {col_idx} ({sheet.cell(row=1, column=col_idx).value}): "
                                         f"Expected {expected}, got {actual}")
            elif isinstance(expected, date):
                if isinstance(actual.date(), date):
                    self.assertEqual(actual.date(), expected,
                                   msg=f"Row {row_idx}, Col {col_idx}: Expected {expected}, got {actual.date()}")
                else:
                    self.fail(f"Row {row_idx}, Col {col_idx}: Expected date {expected}, got {actual} ({actual.__class__.__name__})")
            else:
                self.assertEqual(actual or "", expected or "",
                               msg=f"Row {row_idx}, Col {col_idx} ({sheet.cell(row=1, column=col_idx).value}): "
                                   f"Expected '{expected}', got '{actual}'")

    # --- NOSI Form 1 Tests ---

    def test_export_form1_single_employee(self):
        """Test exporting NOSI Form 1 for a single employee successfully."""
        wizard = self._create_wizard('form1', [self.employee_1.id])
        action = wizard.action_generate_report()
        attachment = self._get_attachment_from_action(action)
        self.assertEqual(attachment.name, "NOSI_Form_1_Export.xlsx")
        sheet = self._read_xlsx_content(attachment)
        expected_form1_headers = [
            "Company Name", "Company NOSI Code", "Social Security No", "Identification Number", "Name",
            "Nationality", "Job Position", "NOSI Registration Start Date", "Basic Salary",
            "Social Insurance Reference Amount", "Total Salary", "Phone Number", "Email", "Address",
        ]
        self._validate_form_headers(sheet, expected_form1_headers)

        addr = self.employee_1.address_id
        address_parts = [addr.street, addr.street2, addr.city, addr.state_id.name, addr.zip, addr.country_id.name]
        expected_address = ", ".join(part for part in address_parts if part)
        # Total Salary = Wage + HOU + TA + OA (5000 + 100 + 110 + 80) = 5290
        expected_form1_results = [self.employee_1.company_id.name, self.employee_1.company_id.l10n_eg_nosi_code,
            self.employee_1.l10n_eg_ssn, self.employee_1.identification_id, self.employee_1.name, self.employee_1.country_id.display_name,
            self.employee_1.job_id.name, self.employee_1.l10n_eg_nosi_registration_start_date, self.employee_1.wage,
            self.employee_1.l10n_eg_social_insurance_reference, 5290, self.employee_1.private_phone,
            self.employee_1.private_email, expected_address,
        ]
        self._validate_employee_row(sheet, 2, self.employee_1, expected_form1_results)

    def test_export_form1_multiple_employees(self):
        """Test exporting NOSI Form 1 for multiple employees successfully."""
        employees = self.employee_1 | self.employee_2
        wizard = self._create_wizard('form1', employees.ids)
        action = wizard.action_generate_report()
        attachment = self._get_attachment_from_action(action)
        self.assertEqual(attachment.name, "NOSI_Form_1_Export.xlsx")

    def test_export_form1_missing_employee_field(self):
        """Test that exporting Form 1 raises an error if an employee field is missing."""
        self.employee_1.l10n_eg_ssn = False  # Set a mandatory field to a falsy value
        wizard = self._create_wizard('form1', [self.employee_1.id])
        with self.assertRaisesRegex(UserError, "have at least one of the following fields missing: Social Security No"):
            wizard.action_generate_report()

    def test_export_form1_missing_company_field(self):
        """Test that exporting Form 1 raises an error if a company field is missing."""
        self.company.l10n_eg_nosi_code = False
        wizard = self._create_wizard('form1', [self.employee_1.id])
        with self.assertRaisesRegex(UserError, "have at least one of the following fields missing: Company NOSI Code"):
            wizard.action_generate_report()

    # --- NOSI Form 6 Tests ---

    def test_export_form6_single_employee(self):
        """Test exporting NOSI Form 6 for a single archived employee successfully."""
        self.env['hr.employee.departure'].create({
            'employee_id': self.employee_1.id,
            'dismissal_date': date(2024, 1, 31),
            'departure_reason_id': self.departure_reason.id,
        }).action_register()
        wizard = self._create_wizard('form6', [self.employee_1.id])
        action = wizard.action_generate_report()
        attachment = self._get_attachment_from_action(action)
        self.assertEqual(attachment.name, "NOSI_Form_6_Export.xlsx")
        sheet = self._read_xlsx_content(attachment)
        expected_form6_headers = [
            "Company Name", "Company NOSI Code", "Social Security No", "Identification Number", "Name",
            "Subscription End Date", "End Reason", "Phone Number", "Email", "Address",
        ]
        self._validate_form_headers(sheet, expected_form6_headers)

        addr = self.employee_1.address_id
        address_parts = [addr.street, addr.street2, addr.city, addr.state_id.name, addr.zip, addr.country_id.name]
        expected_address = ", ".join(part for part in address_parts if part)
        expected_form1_results = [self.employee_1.company_id.name, self.employee_1.company_id.l10n_eg_nosi_code,
            self.employee_1.l10n_eg_ssn, self.employee_1.identification_id, self.employee_1.name, self.employee_1.departure_date,
            self.employee_1.departure_reason_id.name, self.employee_1.private_phone, self.employee_1.private_email, expected_address,
        ]
        self._validate_employee_row(sheet, 2, self.employee_1, expected_form1_results)

    def test_export_form6_multiple_employees(self):
        """Test exporting NOSI Form 6 for multiple archived employees successfully."""
        self.env['hr.employee.departure'].create([
            {
                'employee_id': self.employee_1.id,
                'dismissal_date': date(2024, 2, 28),
                'departure_reason_id': self.departure_reason.id,
            },
            {
                'employee_id': self.employee_2.id,
                'dismissal_date': date(2024, 2, 28),
                'departure_reason_id': self.departure_reason.id,
            },
        ]).action_register()
        wizard = self._create_wizard('form6', (self.employee_1 + self.employee_2).ids)
        action = wizard.action_generate_report()
        attachment = self._get_attachment_from_action(action)
        self.assertEqual(attachment.name, "NOSI_Form_6_Export.xlsx")

    def test_export_form6_missing_employee_field(self):
        """Test that exporting Form 6 raises an error if a mandatory field is missing."""
        self.employee_1.write({
            'active': False,
            'departure_date': False,
            'departure_reason_id': self.departure_reason.id,
        })
        wizard = self._create_wizard('form6', [self.employee_1.id])
        with self.assertRaisesRegex(UserError, "have at least one of the following fields missing: Subscription End Date"):
            wizard.action_generate_report()

    def test_export_form6_for_active_employee(self):
        """Test that exporting Form 6 raises an error for an active employee."""
        self.assertTrue(self.employee_1.active)
        wizard = self._create_wizard('form6', [self.employee_1.id])
        with self.assertRaisesRegex(UserError, "You can only export NOSI Form 6 for archived employees."):
            wizard.action_generate_report()
