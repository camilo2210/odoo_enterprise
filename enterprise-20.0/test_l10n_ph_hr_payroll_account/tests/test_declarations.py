# Part of Odoo. See LICENSE file for full copyright and licensing details.
import io
from datetime import date

from dateutil.relativedelta import relativedelta
from dateutil.rrule import MONTHLY, rrule
from freezegun import freeze_time
from odoo.addons.test_l10n_ph_hr_payroll_account.tests.common import (
    TestL10NPhHrPayrollCommon,
)
from odoo.tests.common import tagged
from odoo.tools import mute_logger

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


@tagged("post_install", "post_install_l10n", "-at_install", "payslips_validation")
class TestDeclarations(TestL10NPhHrPayrollCommon):

    _test_user_groups = None  # FIXME list needed groups

    @classmethod
    @mute_logger('odoo.models.unlink')  # avoid spamming the logs with 'deleted' from the payslip's computation
    def setUpClass(cls):
        super().setUpClass()
        regular_employee_type = cls.env.ref('l10n_ph_hr_payroll.l10n_ph_contract_type_regular')
        cls.high_earner_employee = cls._setup_employee(
            country=cls.country,
            structure_type=cls.structure_type,
            resource_calendar=cls.resource_calendar,
            contract_fields={
                'date_version': date(2021, 1, 1),
                'contract_date_start': date(2021, 1, 1),
                'wage': 75000.0,
                'schedule_pay': 'semi-monthly',
                'l10n_ph_hr_payroll_employee_rank': 'managerial',
                "is_non_resident": True,
                "employee_type_id": regular_employee_type.id,
            },
            employee_fields={
                "private_phone": "0917-888-9999",
                "birthday": "1975-08-15",
                "private_street": "Penthouse 1, The Residences",
                "private_street2": "Greenbelt, Ayala Center",
                "private_city": "Makati City",
                "private_state_id": cls.env.ref("base.state_ph_01").id,
                "private_zip": "1228",
                "private_country_id": cls.country.id,
                "l10n_ph_hr_payroll_rdo_code": "047",
                "l10n_ph_hr_payroll_registered_address": "Penthouse 1, The Residences, Greenbelt, Makati City",
                "l10n_ph_hr_payroll_registered_zip": "1228",
                "l10n_ph_hr_payroll_work_tin": "111-222-333-000",
                "work_location_id": cls.work_location.id,
                "l10n_ph_hr_payroll_foreign_address": "123 Liberty St, New York, NY 10006, USA",
            },
        )
        cls.high_earner_employee.version_id._set_property_input_value("DM_RICE", 2500)
        cls.high_earner_employee.version_id._set_property_input_value("ECOLA", 2000)
        cls.env['hr.employee.departure'].create([{
            'employee_id': cls.high_earner_employee.id,
            'dismissal_date': date(2026, 12, 16),
            'departure_reason_id': cls.env.ref('l10n_ph_hr_payroll.hr_departure_reason_mandatory_retirement').id,
        }])
        cls._set_test_employee(cls.high_earner_employee)
        cls._make_work_entry(cls, 'l10n_ph_hr_payroll_overtime', date(2026, 12, 5), 17, 19)
        cls.minimum_wage_employee = cls._setup_employee(
            country=cls.country,
            structure_type=cls.structure_type,
            resource_calendar=cls.resource_calendar,
            contract_fields={
                'date_version': date(2026, 3, 1),
                'contract_date_start': date(2026, 3, 1),
                'wage': 7000.0,
                'schedule_pay': 'semi-monthly',
                'l10n_ph_hr_payroll_minimum_wage_earner': True,
                "employee_type_id": regular_employee_type.id,
            },
            employee_fields={
                "private_phone": "0917-555-0123",
                "birthday": "1982-03-29",
                "private_street": "Unit 402, Blue Residences",
                "private_street2": "Katipunan Avenue, Brgy. Loyola Heights",
                "private_city": "Quezon City",
                "private_state_id": cls.env.ref("base.state_ph_01").id,
                "private_zip": "1108",
                "private_country_id": cls.country.id,
                "l10n_ph_hr_payroll_rdo_code": "039",
                "l10n_ph_hr_payroll_registered_address": "Unit 402, Blue Residences, Katipunan Ave., Quezon City",
                "l10n_ph_hr_payroll_registered_zip": "1108",
                "l10n_ph_hr_payroll_work_tin": "294-301-482-000",
                "work_location_id": cls.work_location.id,
            },
        )
        cls.env['l10n_ph_hr_payroll.previous_employment'].create({
            "employee_id": cls.minimum_wage_employee.id,
            "version_ids": cls.minimum_wage_employee.version_id.ids,
            "tin": "987-654-321-000",
            "name": "Previous PH Corp",
            "address": "Cebu City",
            "zip": "6000",
            "taxable_basic_salary": 15000.0,
            "taxable_13th_month": 10000.0,
            "taxable_salaries_other": 25000.0,
            "tax_withheld": 2500.0,
            "nontax_basic_mwe": 120000.0,
            "holiday_pay": 2000.0,
            "overtime_pay": 5000.0,
            "night_shift_differential": 1500.0,
            "hazard_pay": 1000.0,
            "nontax_13th_month": 10000.0,
            "nontax_de_minimis": 5000.0,
            "nontax_statutory_contributions": 8000.0,
            "nontax_salaries_other": 2500.0,
        })
        cls._set_test_employee(cls.minimum_wage_employee)
        cls._make_holiday(cls, "l10n_ph_hr_payroll_rh_leave", date(2026, 1, 2))
        cls._make_work_entry(cls, "ph_work_entry_type_attendance", date(2026, 1, 2), 8, 17)
        cls._make_work_entry(cls, "l10n_ph_hr_payroll_ns", date(2026, 1, 6), 21, 24)
        cls._make_work_entry(cls, "l10n_ph_hr_payroll_overtime", date(2026, 1, 8), 17, 19)
        cls._make_work_entry(cls, 'l10n_ph_hr_payroll_overtime', date(2026, 12, 5), 17, 19)
        cls.low_wage_employee = cls._setup_employee(
            country=cls.country,
            structure_type=cls.structure_type,
            resource_calendar=cls.resource_calendar,
            contract_fields={
                'date_version': date(2026, 1, 1),
                'contract_date_start': date(2026, 1, 1),
                'wage': 10000.0,  # Has to fall below the taxable minimum after contributions
                'schedule_pay': 'semi-monthly',
                "employee_type_id": regular_employee_type.id,
            },
            employee_fields={
                "private_phone": "0917-444-5555",
                "birthday": "1990-05-20",
                "private_street": "12 Mabini Street",
                "private_city": "Pasig City",
                "private_state_id": cls.env.ref("base.state_ph_01").id,
                "private_zip": "1600",
                "private_country_id": cls.country.id,
                "l10n_ph_hr_payroll_rdo_code": "043",
                "l10n_ph_hr_payroll_registered_address": "12 Mabini Street, Pasig City",
                "l10n_ph_hr_payroll_registered_zip": "1600",
                "l10n_ph_hr_payroll_work_tin": "222-333-444-000",
                "work_location_id": cls.work_location.id,
            }
        )
        cls._set_test_employee(cls.low_wage_employee)
        cls._make_work_entry(cls, 'l10n_ph_hr_payroll_overtime', date(2026, 12, 5), 17, 19)
        cls.refund_employee = cls._setup_employee(
            country=cls.country,
            structure_type=cls.structure_type,
            resource_calendar=cls.resource_calendar,
            contract_fields={
                'date_version': date(2026, 7, 1),  # Started mid-year
                'contract_date_start': date(2026, 7, 1),
                'wage': 40000.0,
                'schedule_pay': 'semi-monthly',
                "employee_type_id": regular_employee_type.id,
            },
            employee_fields={
                "private_phone": "0917-777-8888",
                "birthday": "1992-11-10",
                "private_street": "789 Pine St",
                "private_city": "Taguig City",
                "private_state_id": cls.env.ref("base.state_ph_01").id,
                "private_zip": "1630",
                "private_country_id": cls.country.id,
                "l10n_ph_hr_payroll_rdo_code": "044",
                "l10n_ph_hr_payroll_registered_address": "789 Pine St, Taguig City",
                "l10n_ph_hr_payroll_registered_zip": "1630",
                "l10n_ph_hr_payroll_work_tin": "333-444-555-000",
                "work_location_id": cls.work_location.id,
            }
        )
        # Previous employment with massive tax withheld to trigger a refund
        cls.env['l10n_ph_hr_payroll.previous_employment'].create({
            "employee_id": cls.refund_employee.id,
            "version_ids": cls.refund_employee.version_id.ids,
            "tin": "444-555-666-000",
            "name": "Over-Taxing Corp",
            "address": "Makati City",
            "zip": "1200",
            "taxable_basic_salary": 200000.0,
            "taxable_13th_month": 15000.0,
            "taxable_salaries_other": 10000.0,
            "tax_withheld": 85000.0,
            "nontax_13th_month": 15000.0,
            "nontax_statutory_contributions": 12000.0,
        })
        # Most of the reports needs a full year of payslips to run. To avoid doing the calculation many times, we will
        # prepare all of these at once.
        payruns_data = []
        payrun_date = date(2026, 1, 1)
        end_date = date(2026, 12, 31)
        employees = cls.high_earner_employee | cls.minimum_wage_employee | cls.low_wage_employee | cls.refund_employee
        for dt in rrule(MONTHLY, dtstart=payrun_date, until=end_date):
            concerned_versions = employees.version_id.filtered(lambda v: v._is_in_contract(dt.date()))
            for day_start, day_end in ((0, 15), (16, 31)):
                payruns_data.append({
                    'date_start': dt.date() + relativedelta(day=day_start),
                    'date_end': dt.date() + relativedelta(day=day_end),
                    'structure_id': employees.structure_type_id.default_struct_id.id,
                    'version_ids': concerned_versions.ids,
                })

        payruns = cls.env['hr.payslip.run'].create(payruns_data)
        for i, payrun in enumerate(payruns):
            payslips = payrun._generate_payslips()
            for payslip in payslips:
                month = payrun.date_start.month
                day = payrun.date_start.day
                if month == 3 and day == 1:
                    payslip.version_id._set_property_input_value("DM_RICE", 2500)
                if month == 6 and day == 1:
                    payslip._set_input_value("NCFB", 50000)
                    payslip._set_input_value("CA", 5000)
                if month == 12 and day == 1:
                    payslip._set_input_value('CA', 110_000)
                if month == 12 and day == 16:
                    payslip._set_input_value("DM_CLOTHING", 100_000)
                    payslip.version_id._set_property_input_value('DM_RICE', 1000)
                    payslip._set_input_value('DM_RICE', 1000)
                    payslip.compute_sheet()
            payrun.action_validate()
        payruns.action_paid()

    @freeze_time('2027-01-2')
    def test_1601c(self):
        """ Test the generation of the 1601c and ensure that the amounts are correct. """
        form = self.env['l10n_ph_hr_payroll.form_1601c'].create({})
        form.action_generate_declarations()
        form.action_process_declaration()

        form.write({
            'tax_previously_remitted': 100.0,
            'other_remittances': 25.0,
            'surcharge': 25.0,
            'interest': 50.0,
            'compromise': 75.0,
        })

        self.assertRecordValues(
            form,
            [{
                'mwe_basic': 14000.0,  # 7000 * 2 => 14000
                'mwe_holiday_ot': 194.47,  # 194.47 * 1 => 194.47
                'non_taxable_benefits': 262500.0,
                'de_minimis': 40500.0,
                'mandatory_contributions': 13150.0,
                'other_non_taxable': 665753.42,  # 665753.42 * 1 => 665753.42
                'total_non_taxable': 996097.89,
                'total_taxable': 936815.12,
                'taxable_exempt': 0.0,
                'taxable_taxed': 936815.12,
                'total_tax_withheld': 218529.66,
                'total_tax_adjustment': -152347.42,
                'total_tax_withheld_for_rem': 66182.24,
                'tax_previously_remitted': 100.0,
                'other_remittances': 25.0,
                'total_tax_remittances': 125.0,  # 100 + 25 => 125
                'tax_still_due': 66057.24,
                'surcharge': 25.0,
                'interest': 50.0,
                'compromise': 75.0,
                'penalties': 150.0,  # 25 + 50 + 75 => 150
                'grand_total_compensation': 1932913.01,
                'total_still_due': 66207.24,
            }]
        )

    @freeze_time('2027-01-1')
    def test_2316_minimum_wage(self):
        """ Test the generation of the 2316 and ensure that the amounts are correct. """
        declaration = self.env['l10n_ph_hr_payroll.form_2316'].create({})
        rendering_data = declaration._get_rendering_data(self.minimum_wage_employee)
        self.assertDictEqual(
            rendering_data,
            {
                "employees_with_error": {},
                self.minimum_wage_employee: {
                    "employee_branch_code": "000",
                    "employee_contact_number": "09175550123",
                    "employee_dob": "03/29/1982",
                    "employee_first_name": "PH",
                    "employee_foreign_address": 'False',
                    "employee_formatted_name": "Employee, PH",
                    "employee_is_mwe": 'True',
                    "employee_last_name": "Employee",
                    "employee_legal_name": "PH Employee",
                    "employee_local_home_address": "Unit 402, Blue Residences, Katipunan Avenue, Brgy. Loyola Heights, Quezon City, National Capital Region",
                    "employee_local_home_zip": "1108",
                    "employee_middle_name": 'False',
                    "employee_rdo_code": "039",
                    "employee_registered_address": "Unit 402, Blue Residences, Katipunan Ave., Quezon City",
                    "employee_registered_zip": "1108",
                    "employee_stat_min_wage_day": "695.00",
                    "employee_stat_min_wage_month": "21139.58",
                    "employee_tin": "294301482",
                    "employer_branch_code": "000",
                    "employer_is_main": 'True',
                    "employer_name": "My Philippines Company Inc.",
                    "employer_previous_address": "Cebu City",
                    "employer_previous_name": "Previous PH Corp",
                    "employer_previous_tin": "987654321000",
                    "employer_previous_zip": "6000",
                    "employer_registered_address": "12th Floor, Tower 1, Ayala Triangle, Ayala Avenue, Brgy. Bel-Air, Makati City, National Capital Region",
                    "employer_signatory_name": "My Philippines Company Inc.",
                    "employer_tin": "123456789",
                    "employer_zip": "1226",
                    "item_19": "465361.07",
                    "item_20": "253194.47",  # Reduced by 10k
                    "item_21": "212166.60",  # Increased by 10k
                    "item_22": "50000.00",
                    "item_23": "262166.60",  # Increased by 10k
                    "item_24": "1824.99",    # Increased by 1500 tax
                    "item_25a": "0.00",
                    "item_25b": "2500.00",
                    "item_26": "2500.00",
                    "item_27": "0.00",
                    "item_28": "2500.00",
                    "item_29": "126200.00",
                    "item_30": "34.52",
                    "item_31": "159.95",
                    "item_32": "0.00",
                    "item_33": "0.00",
                    "item_34": "80000.00",   # Reduced by 10k (shifted to taxable)
                    "item_35": "33000.00",
                    "item_36": "13800.00",
                    "item_37": "0.00",
                    "item_38": "253194.47",  # Reduced by 10k
                    "item_39": "0.00",
                    "item_40": "",
                    "item_41": "",
                    "item_42": "",
                    "item_43": "",
                    "item_44a": "115000.00",
                    "item_44a_label": "Taxable Allowances",
                    "item_44b": "50000.00",
                    "item_44b_label": "Non-Cash Earnings",
                    "item_45": "",
                    "item_46": "",
                    "item_47": "",
                    "item_48": "47166.60",   # Increased by 10k
                    "item_49": "",
                    "item_50": "",
                    "item_51a": "0.00",
                    "item_51a_label": "",
                    "item_51b": "0.00",
                    "item_51b_label": "",
                    "item_52": "212166.60",  # Increased by 10k
                    "period_from": "03/01",
                    "period_to": "12/31",
                    "year": "2026",
                },
            },
        )

    @freeze_time("2027-01-1")
    def test_2316_high_earner(self):
        """ Test the generation of the 2316 for a high earner / managerial employee to ensure heavy tax and FBT rules. """
        declaration = self.env["l10n_ph_hr_payroll.form_2316"].create({})
        rendering_data = declaration._get_rendering_data(self.high_earner_employee)
        self.assertDictEqual(
            rendering_data,
            {
                "employees_with_error": {},
                self.high_earner_employee: {
                    "employee_branch_code": "000",
                    "employee_contact_number": "09178889999",
                    "employee_dob": "08/15/1975",
                    "employee_first_name": "PH",
                    "employee_foreign_address": "123 Liberty St, New York, NY 10006, USA",
                    "employee_formatted_name": "Employee, PH",
                    "employee_is_mwe": 'False',
                    "employee_last_name": "Employee",
                    "employee_legal_name": "PH Employee",
                    "employee_local_home_address": "Penthouse 1, The Residences, Greenbelt, Ayala Center, Makati City, National Capital Region",
                    "employee_local_home_zip": "1228",
                    "employee_middle_name": 'False',
                    "employee_rdo_code": "047",
                    "employee_registered_address": "Penthouse 1, The Residences, Greenbelt, Makati City",
                    "employee_registered_zip": "1228",
                    "employee_stat_min_wage_day": "0.00",
                    "employee_stat_min_wage_month": "0.00",
                    "employee_tin": "111222333",
                    "employer_branch_code": "000",
                    "employer_is_main": 'True',
                    "employer_name": "My Philippines Company Inc.",
                    "employer_previous_address": "",
                    "employer_previous_name": "",
                    "employer_previous_tin": "",
                    "employer_previous_zip": "",
                    "employer_registered_address": "12th Floor, Tower 1, Ayala Triangle, Ayala Avenue, Brgy. Bel-Air, Makati City, National Capital Region",
                    "employer_signatory_name": "My Philippines Company Inc.",
                    "employer_tin": "123456789",
                    "employer_zip": "1226",
                    "item_19": "2862774.25",
                    "item_20": "847153.42",
                    "item_21": "2015620.83",
                    "item_22": "0.00",
                    "item_23": "2015620.83",
                    "item_24": "407186.25",
                    "item_25a": "407186.25",
                    "item_25b": "0.00",
                    "item_26": "407186.25",
                    "item_27": "0.00",
                    "item_28": "407186.25",
                    "item_29": "0.00",
                    "item_30": "0.00",
                    "item_31": "0.00",
                    "item_32": "0.00",
                    "item_33": "0.00",
                    "item_34": "90000.00",
                    "item_35": "38000.00",
                    "item_36": "53400.00",
                    "item_37": "665753.42",
                    "item_38": "847153.42",
                    "item_39": "1677850.00",
                    "item_40": "",
                    "item_41": "",
                    "item_42": "48000.00",
                    "item_43": "",
                    "item_44a": "115000.00",
                    "item_44a_label": "Taxable Allowances",
                    "item_44b": "0.00",
                    "item_44b_label": "",
                    "item_45": "",
                    "item_46": "",
                    "item_47": "",
                    "item_48": "174770.83",
                    "item_49": "",
                    "item_50": "",
                    "item_51a": "0.00",
                    "item_51a_label": "",
                    "item_51b": "0.00",
                    "item_51b_label": "",
                    "item_52": "2015620.83",
                    "period_from": "01/01",
                    "period_to": "12/16",
                    "year": "2026",
                },
            },
        )

    @freeze_time("2027-01-1")
    def test_2316_populate_departing_vs_all(self):
        """ Test the populate function of the 2316 to ensure it correctly filters employees. """
        declaration = self.env["l10n_ph_hr_payroll.form_2316"].create({})
        # 1. Only departing employee
        declaration.departing_employees_only = True
        declaration.action_generate_declarations()
        self.assertEqual(declaration.lines_count, 1)
        # 2. Everyone
        declaration.departing_employees_only = False
        declaration.action_generate_declarations()
        self.assertEqual(declaration.lines_count, 4)

    @freeze_time("2027-01-1")
    def test_2316_low_wage(self):
        """ A regular (non-MWE) employee whose yearly basic salary is under the 250k tax-exempt cap
            must have their basic salary reported as non-taxable (item 29), not taxable (item 39). """
        declaration = self.env["l10n_ph_hr_payroll.form_2316"].create({})
        rendering_data = declaration._get_rendering_data(self.low_wage_employee)
        self.assertDictEqual(
            rendering_data,
            {
                "employees_with_error": {},
                self.low_wage_employee: {
                    "employee_branch_code": "000",
                    "employee_contact_number": "09174445555",
                    "employee_dob": "05/20/1990",
                    "employee_first_name": "PH",
                    "employee_foreign_address": "False",
                    "employee_formatted_name": "Employee, PH",
                    "employee_is_mwe": "False",
                    "employee_last_name": "Employee",
                    "employee_legal_name": "PH Employee",
                    "employee_local_home_address": "12 Mabini Street, Pasig City, National Capital Region",
                    "employee_local_home_zip": "1600",
                    "employee_middle_name": "False",
                    "employee_rdo_code": "043",
                    "employee_registered_address": "12 Mabini Street, Pasig City",
                    "employee_registered_zip": "1600",
                    "employee_stat_min_wage_day": "0.00",
                    "employee_stat_min_wage_month": "0.00",
                    "employee_tin": "222333444",
                    "employer_branch_code": "000",
                    "employer_is_main": "True",
                    "employer_name": "My Philippines Company Inc.",
                    "employer_previous_address": "",
                    "employer_previous_name": "",
                    "employer_previous_tin": "",
                    "employer_previous_zip": "",
                    "employer_registered_address": "12th Floor, Tower 1, Ayala Triangle, Ayala Avenue, Brgy. Bel-Air, Makati City, National Capital Region",
                    "employer_signatory_name": "My Philippines Company Inc.",
                    "employer_tin": "123456789",
                    "employer_zip": "1226",
                    "item_19": "573777.73",
                    "item_20": "363000.00",
                    "item_21": "210777.73",
                    "item_22": "0.00",
                    "item_23": "210777.73",
                    "item_24": "0.00",
                    "item_25a": "28375.55",
                    "item_25b": "0.00",
                    "item_26": "28375.55",
                    "item_27": "0.00",
                    "item_28": "28375.55",
                    "item_29": "218600.00",
                    "item_30": "0.00",
                    "item_31": "0.00",
                    "item_32": "0.00",
                    "item_33": "0.00",
                    "item_34": "90000.00",
                    "item_35": "33000.00",
                    "item_36": "21400.00",
                    "item_37": "0.00",
                    "item_38": "363000.00",
                    "item_39": "0.00",
                    "item_40": "",
                    "item_41": "",
                    "item_42": "",
                    "item_43": "",
                    "item_44a": "115000.00",
                    "item_44a_label": "Taxable Allowances",
                    "item_44b": "50000.00",
                    "item_44b_label": "Non-Cash Earnings",
                    "item_45": "",
                    "item_46": "",
                    "item_47": "",
                    "item_48": "45499.92",
                    "item_49": "",
                    "item_50": "228.49",
                    "item_51a": "49.32",
                    "item_51a_label": "Holiday Pay",
                    "item_51b": "0.00",
                    "item_51b_label": "",
                    "item_52": "210777.73",
                    "period_from": "01/01",
                    "period_to": "12/31",
                    "year": "2026",
                },
            },
        )

    # /!\ If possible, always assert changes in test added for 1604-C (the dat file) using the alphalist validation
    # software to ensure that everything is as expected.

    @freeze_time("2027-01-1")
    def test_1604c(self):
        """ Test generating the 1604-C for all of our test employees. """
        self.env.company.vat = '123-456-789-00001'
        # Create one year of 1601-C
        for dt in rrule(MONTHLY, dtstart=date(2026, 1, 1), until=date(2026, 12, 31)):
            form = self.env['l10n_ph_hr_payroll.form_1601c'].create({
                'period_start_date': dt,
                'period_end_date': dt + relativedelta(day=31),
                'remittance_date': dt + relativedelta(day=28),
            })
            form.action_generate_declarations()
            form.action_process_declaration()
            form.action_confirm_declaration()

        # Create the yearly 2316 for all employees.
        declaration = self.env["l10n_ph_hr_payroll.form_2316"].create({})
        declaration.action_generate_declarations()
        declaration.action_confirm_declaration()

        # Finally, create the 1604-C
        declaration = self.env["l10n_ph_hr_payroll.form_1604c"].create({})
        declaration.action_generate_declarations()
        summary_file, dat_file = declaration.action_generate_files()
        # Validate the summary data
        if load_workbook is not None:
            expected_values = {
                0: ('Remittance per BIR Form No. 1601-C', None, None, None, None, None, None),
                1: ('MONTH', 'DATE OF REMITTANCE', 'NAME OF BANK/BANK CODE/ROR NO., IF ANY', 'TAXES WITHHELD', 'ADJUSTMENT', 'PENALTIES', 'TOTAL AMOUNT REMITTED'),
                2: ('JAN',   '28-Jan-2026',         None,                                    29262.4,          0,             0,           29262.4),
                3: ('FEB',   '28-Feb-2026',         None,                                    29262.4,          0,             0,           29262.4),
                4: ('MAR',   '28-Mar-2026',         None,                                    29262.4,          0,             0,           29262.4),
                5: ('APR',   '28-Apr-2026',         None,                                    29262.4,          0,             0,           29262.4),
                6: ('MAY',   '28-May-2026',         None,                                    29262.4,          0,             0,           29262.4),
                7: ('JUN',   '28-Jun-2026',         None,                                    51893.55,         0,             0,           51893.55),
                8: ('JUL',   '28-Jul-2026',         None,                                    40149.8,          0,             0,           40149.8),
                9: ('AUG',   '28-Aug-2026',         None,                                    40149.8,          0,             0,           40149.8),
                10: ('SEP',  '28-Sep-2026',         None,                                    40149.8,          0,             0,           40149.8),
                11: ('OCT',  '28-Oct-2026',         None,                                    40149.8,          0,             0,           40149.8),
                12: ('NOV',  '28-Nov-2026',         None,                                    40149.8,          0,             0,           40149.8),
                13: ('DEC',  '28-Dec-2026',         None,                                    218529.66,        -152347.42,    0,           66182.24),
                14: (None,   None,                  None,                                    617484.21,        -152347.42,    0,           465136.79),
            }

            summary_file_content = io.BytesIO(summary_file.raw.content)
            xlsx = load_workbook(filename=summary_file_content, data_only=True)
            sheet = xlsx.worksheets[0]
            sheet_values = dict(enumerate(sheet.values))
            self.assertDictEqual(expected_values, sheet_values)

        # Validate the dat data
        expected_values = {
            0: ['H1604C', '123456789', '0001', '12/31/2026'],
            1: ['D1', '1604C', '123456789', '0001', '12/31/2026', '1', '111222333', '0000', 'Employee', 'PH', '', 'NCR', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '01/01/2026', '12/16/2026', '2862774.25', '0.00', '90000.00', '38000.00', '53400.00', '665753.42', '847153.42', '1677850.00', '174770.83', '163000.00', '2015620.83', '2015620.83', '2015620.83', '407186.25', '0.00', '370339.70', '36846.55', '0.00', '407186.25', 'PH', 'R', 'TR', 'No', '0.00'],
            2: ['D1', '1604C', '123456789', '0001', '12/31/2026', '2', '222333444', '0000', 'Employee', 'PH', '', 'NCR', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '0.00', '01/01/2026', '12/31/2026', '573777.73', '218600.00', '90000.00', '33000.00', '21400.00', '0.00', '363000.00', '0.00', '45499.92', '165277.81', '210777.73', '210777.73', '210777.73', '0.00', '0.00', '39399.09', '0.00', '39399.09', '0.00', 'PH', 'R', 'NA', 'No', '0.00'],
            3: ['D1', '1604C', '123456789', '0001', '12/31/2026', '3', '333444555', '0000', 'Employee', 'PH', '', 'NCR', '252000.00', '0.00', '15000.00', '0.00', '12000.00', '0.00', '27000.00', '200000.00', '15000.00', '10000.00', '225000.00', '07/01/2026', '12/31/2026', '730999.96', '0.00', '75000.00', '9000.00', '23700.00', '0.00', '107700.00', '456300.00', '56999.96', '110000.00', '623299.96', '848299.96', '848299.96', '114574.99', '85000.00', '90615.30', '0.00', '61040.31', '114574.99', 'PH', 'R', 'NA', 'No', '0.00'],
            4: ['C1', '1604C', '123456789', '0001', '12/31/2026', '252000.00', '0.00', '15000.00', '0.00', '12000.00', '0.00', '27000.00', '200000.00', '15000.00', '10000.00', '225000.00', '4167551.94', '218600.00', '255000.00', '80000.00', '98500.00', '665753.42', '1317853.42', '2134150.00', '277270.71', '438277.81', '2849698.52', '3074698.52', '3074698.52', '521761.24', '85000.00', '500354.09', '36846.55', '100439.40', '521761.24', '0.00'],
            5: ['D2', '1604C', '123456789', '0001', '12/31/2026', '1', '294301482', '0000', 'Employee', 'PH', '', 'NCR', '205000.00', '120000.00', '2000.00', '5000.00', '1500.00', '1000.00', '10000.00', '5000.00', '8000.00', '2500.00', '155000.00', '10000.00', '40000.00', '50000.00', '03/01/2026', '12/31/2026', '451561.07', '695.00', '21139.58', '21139.58', '365.00', '34.52', '159.95', '0.00', '0.00', '80000.00', '33000.00', '13800.00', '0.00', '239394.47', '47166.60', '165000.00', '212166.60', '262166.60', '262166.60', '1824.99', '2500.00', '33894.50', '0.00', '34569.51', '1824.99', 'PH', 'R', 'NA', 'No', '0.00', '112400.00'],
            6: ['C2', '1604C', '123456789', '0001', '12/31/2026', '205000.00', '120000.00', '2000.00', '5000.00', '1500.00', '1000.00', '10000.00', '5000.00', '8000.00', '2500.00', '155000.00', '10000.00', '40000.00', '50000.00', '451561.07', '695.00', '21139.58', '21139.58', '34.52', '159.95', '0.00', '0.00', '80000.00', '33000.00', '13800.00', '0.00', '239394.47', '47166.60', '165000.00', '212166.60', '262166.60', '262166.60', '1824.99', '2500.00', '33894.50', '0.00', '34569.51', '1824.99', '0.00', '112400.00'],
        }
        dat_file_content = dat_file.raw.content.decode()
        dat_file_values = dict(enumerate([vals.split(',') for vals in dat_file_content.split('\n')]))
        self.assertDictEqual(expected_values, dat_file_values)
