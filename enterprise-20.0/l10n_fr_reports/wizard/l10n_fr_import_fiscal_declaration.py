import io
import xlsxwriter

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from odoo import fields, models
from odoo.exceptions import UserError
from odoo.tools import BinaryBytes

HEADERS = [
    'Line code',
    'Line name',
    'Expression label',
    'Expected format',
    'External Value',
    'Comment',
]


class ImportFiscalDeclaration(models.TransientModel):
    _name = 'l10n_fr.import.fiscal.declaration'
    _description = "Import Fiscal Declaration Wizard"

    data = fields.Binary(string="File")
    filename = fields.Char(string="File Name")
    report_id = fields.Many2one(string="Report", comodel_name='account.report', index='btree')

    def action_download_template(self):
        self.ensure_one()

        section_report_ids = set(
            dict(
                self.env['account.report.expression']._read_group(
                    domain=[
                        ('engine', 'in', ('reference', 'external')),
                        ('report_line_id.report_id', 'in', self.report_id.section_report_ids.ids),
                    ],
                    groupby=['report_line_id.report_id'],
                    aggregates=['id:array_agg'],
                )
            )
        )
        xlsx_data = None
        with io.BytesIO() as output:
            with xlsxwriter.Workbook(output, {'in_memory': True}) as workbook:
                for section_report in section_report_ids:
                    sheet = workbook.add_worksheet(section_report.name)
                    for col, header in enumerate(HEADERS):
                        sheet.write(0, col, header)

                    report_options = section_report.get_options({})
                    # We need to remove the add new section line
                    report_lines = [line for line in section_report._get_lines(report_options) if line.code != 'add_new_section']

                    # Keep track of the row to add, as report lines may use an expression engine other than reference or external.
                    # We cannot rely on enumerate for this reason.
                    row = 1
                    for line, line_data in zip(section_report.line_ids, report_lines):
                        if not line.expression_ids:
                            continue

                        for index, expression in enumerate(line.expression_ids):
                            if expression.engine not in {"reference", "external"}:
                                continue

                            sheet.write(row, 0, line.code)
                            sheet.write(row, 1, line.name)
                            sheet.write(row, 2, expression.label)
                            sheet.write(row, 3, expression.figure_type or 'float')

                            col = next(
                                (col for col in line_data.columns if col.expression_label == expression.label),
                                None,
                            )
                            sheet.write(row, 4, col and col.no_format or '')
                            if expression.figure_type == 'many2one':
                                sheet.write(row, 5, self.env._('This line should be a reference to an existing record, just put the id of that record'))
                            if expression.figure_type == 'boolean':
                                sheet.write(row, 5, self.env._('Write 1 if true, 0 otherwise'))
                            else:
                                sheet.write(row, 5, '')
                            row += 1

            xlsx_data = output.getvalue()
        attachment = self.env['ir.attachment'].create({
            'name': self.env._('template_fiscal_declaration.xlsx'),
            'raw': BinaryBytes(xlsx_data),
            'type': 'binary',
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'new',
        }

    def action_import_template(self):
        self.ensure_one()

        if not self.data:
            raise UserError(self.env._("Please upload a file."))

        report_file = io.BytesIO(self.data)
        workbook = load_workbook(filename=report_file, data_only=True)

        skipped_sheet = []
        worksheets = workbook.worksheets
        for worksheet in worksheets:
            section_report = self.report_id.section_report_ids.filtered(
                lambda r: r.name == worksheet.title
            )[:1]
            if not section_report:
                skipped_sheet.append(section_report)
                continue

            report_options = section_report.get_options({})
            all_column_groups_expression_totals = section_report._compute_expression_totals_for_each_column_group(section_report.line_ids.expression_ids, report_options)
            # Convert all_column_groups_expression_totals to a json-friendly form (its keys are records)
            json_friendly_column_group_totals = section_report._get_json_friendly_column_group_totals(all_column_groups_expression_totals)

            for row_index, cols in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if len(cols) != 6 or (row_index == 1 and len([valid_col for valid_col in cols if valid_col in HEADERS]) != 6):
                    raise UserError(self.env._("Invalid XLSX file (in worksheet %s) , follow the given template", worksheet.title))

                line_code = cols[0]
                expression_label = cols[2]
                external_value = cols[4]
                line = section_report.line_ids.filtered(
                    lambda line: line.code == line_code
                )[:1]
                # No line with that code was found, go to the next one we don't want to block the whole import
                if not line:
                    continue

                expression = line.expression_ids.filtered(
                    lambda expr: expr.label == expression_label
                )[:1]

                if expression.engine not in {'external', 'reference'}:
                    continue

                try:
                    # Some expressions don't have any figure type. For those, they will be interpreted as float
                    if expression.figure_type in {False, 'float', 'monetary'}:
                        value = float(external_value or 0.0)
                    elif expression.figure_type == 'string':
                        value = str(external_value or '')
                    elif expression.figure_type == 'many2one':
                        value = f'{expression.formula}:{self.env[expression.formula].browse(external_value)}' if external_value else False
                    else:
                        value = int(external_value or 0)
                except ValueError:
                    raise UserError(self.env._("The value (%(external_value)s) entered for line: %(line_code)s doesn't have the valid format", external_value=external_value, line_code=line_code))

                section_report.action_modify_manual_value(
                    line_id=line.id,
                    options=report_options,
                    column_group_index=report_options['columns'][-1]['column_group_index'],
                    new_value_str=value,
                    target_expression_id=expression.id,
                    rounding=0,  # We have a rounding 0 on all editable formula
                    json_friendly_column_group_totals=json_friendly_column_group_totals,
                )

        if len(skipped_sheet) == len(worksheets):
            raise UserError(self.env._("All the sheets were skipped, please upload a valid file following the template."))

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': self.env._('Import completed'),
                'type': 'success',
                'sticky': False,
                'next': {
                    'type': 'ir.actions.client',
                    'tag': 'reload',
                },
            },
        }
