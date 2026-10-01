import io
import unittest

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from psycopg2.extras import Json

from odoo.exceptions import RedirectWarning
from odoo.tests import tagged
from odoo.tools import BinaryBytes, file_open
from odoo.addons.l10n_fr_reports.tests.common import TestL10nFrReportsCommon


REPORT_PERIOD_START_DATE = '2025-01-01'
REPORT_PERIOD_END_DATE = '2025-12-31'

REPORTS_TO_EXPORT = {
    '2031': {
        'reports': ['l10n_fr_2031'],
        'handler': 'l10n_fr_reports.2031.report.handler',
    },
    '2031BIS': {
        'reports': ['l10n_fr_2031_annexes'],
        'handler': 'l10n_fr_reports.2031.annexes.report.handler',
    },
    '2033A': {
        'reports': ['l10n_fr_2033_A'],
        'handler': 'l10n_fr_reports.2033.a.report.handler',
    },
    '2033B': {
        'reports': ['l10n_fr_2033_B'],
        'handler': 'l10n_fr_reports.2033.b.report.handler',
    },
    '2033C': {
        'reports': ['l10n_fr_2033_C_1', 'l10n_fr_2033_C_2'],
        'handler': 'l10n_fr_reports.2033.c.report.handler',
    },
    '2033D': {
        'reports': ['l10n_fr_2033_D_1', 'l10n_fr_2033_D_2'],
        'handler': 'l10n_fr_reports.2033.d.report.handler',
    },
    '2033E': {
        'reports': ['l10n_fr_2033_E'],
        'handler': 'l10n_fr_reports.2033.e.report.handler',
    },
    '2033F': {
        'reports': ['l10n_fr_2033_F'],
        'handler': 'l10n_fr_reports.2033.f.report.handler',
    },
    '2033G': {
        'reports': ['l10n_fr_2033_G'],
        'handler': 'l10n_fr_reports.2033.g.report.handler',
    },
    '2065': {
        'reports': ['l10n_fr_2065_SD'],
        'handler': 'l10n_fr_reports.2065_sd.report.handler',
    },
    '2065BIS': {
        'reports': ['l10n_fr_2065_bis_SD'],
        'handler': 'l10n_fr_reports.2065_sd.bis.report.handler',
    },
    '2069RCI': {
        'reports': ['l10n_fr_2069_RCI', 'l10n_fr_2069_RCI_donations'],
        'handler': 'l10n_fr_reports.2069.rci.report.handler',
    },
}


@unittest.skipIf(load_workbook is None, "openpyxl not available")
@tagged('post_install_l10n', 'post_install', '-at_install')
class TestLiasseFiscale(TestL10nFrReportsCommon):
    _test_user_groups = None  # FIXME list needed groups

    @classmethod
    def setUpClass(cls):
        super().setUpClass()

        cls.report = cls.env.ref('l10n_fr_reports.fiscal_declaration')
        cls.handler = cls.env['l10n_fr_reports.fiscal_declaration_handler']

    def _generate_fiscal_options(self, report=False, date_from=REPORT_PERIOD_START_DATE, date_to=REPORT_PERIOD_END_DATE):
        if not report:
            report = self.report
        options = super()._generate_options(report, date_from, date_to)
        options.update({
            'liasse_fiscale': {
                '2033A': True,
                '2033B': True,
                '2033C': True,
                '2033D': True,
                '2033E': True,
                '2033F': True,
                '2033G': True,
                '2031': True,
                '2031BIS': True,
                '2065': True,
                '2065BIS': True,
                '2069RCI': True,
                'special_circ': 'NOR',
                'tax_type': 'IS',
            }
        })
        return options

    def _add_section(self, section, section_options, config):
        for line in section._get_lines(section_options):
            if line.code == 'add_new_section':
                line_dict = {
                    'id': line.id,
                    'parent_id': line.parent_id,
                    'name': line.name,
                }
                self.env[config['handler']].action_add_new_section(section_options, {'line': line_dict})

    def _remove_section(self, lines, section, section_options, config, match_name=None):
        for line in lines:
            for section_line in section._get_lines(section_options):
                if line.code == section_line.code and (not match_name or match_name == line.name):
                    line_dict = {
                        'id': section_line.id,
                        'parent_id': section_line.parent_id,
                        'name': section_line.name,
                    }
                    self.env[config['handler']].action_remove_section(section_options, {'line': line_dict})

    def _set_invalid_siret(self, partner):
        # Store a malformed SIRET directly in the JSON column, bypassing the
        # additional_identifiers constraint (which would otherwise reject it).
        self.env.cr.execute(
            "UPDATE res_partner SET additional_identifiers = %s WHERE id = %s",
            [Json({'FR_SIRET': '1A2B3C4D5E6F7G'}), partner.id],
        )
        # Invalidate the whole cache: the malformed value must propagate to the
        # dependent computed `l10n_fr_siret` and the company's related mirror of it.
        self.env.invalidate_all()

    def test_check_values_export_ok(self):
        options = self._generate_options()
        self.handler._check_values_export(options, self.report)

    def test_check_values_export_ok_no_writer(self):
        # no account_representative => company is the writer
        self.company.account_representative_id = False
        options = self._generate_options()
        self.handler._check_values_export(options, self.report)

    ###############################################################
    # Writer fields errors
    ###############################################################

    def test_check_values_export_writer_missing_siret(self):
        self.firm.l10n_fr_siret = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_writer_invalid_siret(self):
        # Bypass the additional_identifiers constraint to simulate legacy/imported
        # data so that the report's own SIRET validity check is exercised.
        self._set_invalid_siret(self.firm)
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_writer_missing_street(self):
        self.firm.street = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_writer_missing_zip(self):
        self.firm.zip = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_writer_missing_city(self):
        self.firm.city = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_writer_missing_country(self):
        self.firm.country_id = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    ###############################################################
    # Debtor fields errors
    ###############################################################

    def test_check_values_export_debtor_missing_siret(self):
        self.company.partner_id.additional_identifiers = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_debtor_invalid_siret(self):
        # Bypass the additional_identifiers constraint to simulate legacy/imported
        # data so that the report's own SIRET validity check is exercised.
        self._set_invalid_siret(self.company.partner_id)
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_debtor_missing_street(self):
        self.company.partner_id.street = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_debtor_missing_zip(self):
        self.company.partner_id.zip = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_debtor_missing_city(self):
        self.company.partner_id.city = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    def test_check_values_export_debtor_missing_country(self):
        self.company.partner_id.country_id = False
        with self.assertRaises(RedirectWarning):
            self.handler._check_values_export(self._generate_options(), self.report)

    ###############################################################
    # Fiscale reports export
    ###############################################################

    def test_fiscal_export(self):
        expected_file = file_open('l10n_fr_reports/tests/expected_files/fiscal_reports/Fiscal_Declaration_company_1_data_2025.xml', 'rb').read()
        self._generate_move(amount=3000, move_type='out_invoice').action_post()
        self._generate_move(partner=self.partner_b, amount=2500).action_post()
        self._generate_move(partner=self.partner_b, amount=2000, move_type='out_invoice', account=self.copyright_account).action_post()
        self._generate_move(partner=self.partner_a, amount=5000, account=self.fees_account).action_post()

        options = self._generate_fiscal_options()

        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        xml_file = self.env['l10n_fr_reports.fiscal_declaration_handler']._export_fiscal_declaration(options)
        self.assertEqual(xml_file['file_name'], 'Fiscal_Declaration_company_1_data_2025.xml')
        self.assertEqual(xml_file['file_content'], expected_file)

    ###############################################################
    # Fiscale reports extraction
    ###############################################################

    def test_data_extraction_empty(self):
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration_options = options.get('liasse_fiscale')
        for report_name, config in REPORTS_TO_EXPORT.items():
            if not fiscal_declaration_options.get(report_name):
                continue
            lines = []
            mapper = self.env[config['handler']]._get_code_to_edi_id()
            vals = self.env['l10n_fr_reports.fiscal_declaration_handler']._export_report_data(lines, mapper)
            with self.subTest(vals=vals):
                self.assertEqual(vals, [])

    def test_data_extraction_fixed_line(self):
        self._generate_move(amount=3000, move_type='out_invoice').action_post()
        self._generate_move(partner=self.partner_b, amount=2500).action_post()
        self._generate_move(partner=self.partner_b, amount=2000, move_type='out_invoice', account=self.copyright_account).action_post()
        self._generate_move(partner=self.partner_a, amount=5000, account=self.fees_account).action_post()
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        expected = {
            '2031': [
                {'id': 'AA', 'value': ''}, {'id': 'AD', 'value': ''}, {'id': 'KY', 'value': ''},
                {'id': 'AB', 'ftx_1': ''}, {'id': 'BT', 'value': ''}, {'id': 'CA', 'value': 0.0},
                {'id': 'CB', 'value': 0.0}, {'id': 'CC', 'value': 0.0}, {'id': 'CD', 'value': 0.0},
                {'id': 'CE', 'value': 0}, {'id': 'CF', 'value': 0.0}, {'id': 'CH', 'value': 0},
                {'id': 'CJ', 'value': 0}, {'id': 'CK', 'value': 0}, {'id': 'CL', 'value': 0},
                {'id': 'CM', 'value': 0}, {'id': 'AQ', 'value': 0.0}, {'id': 'AE', 'value': 0.0},
                {'id': 'AF', 'value': 0.0}, {'id': 'CP', 'value': 0.0}, {'id': 'KG', 'value': 0.0},
                {'id': 'CN', 'value': 0.0}, {'id': 'CV', 'value': 0.0}, {'id': 'KM', 'value': 0.0},
                {'id': 'AP', 'value': 0.0}, {'id': 'JA', 'value': ''}, {'id': 'KX', 'value': ''},
                {'id': 'KQ', 'value': ''}, {'id': 'JL', 'value': ''}, {'id': 'JB', 'value': ''},
                {'id': 'KR', 'value': ''}, {'id': 'KK', 'value': ''}, {'id': 'AG', 'value': ''},
                {'id': 'AR', 'value': ''}, {'id': 'AY', 'value': ''}, {'id': 'CQ', 'value': 0.0},
                {'id': 'CR', 'value': 0.0}, {'id': 'KZ', 'value': ''}, {'id': 'CW', 'value': 0},
                {'id': 'CX', 'value': 0}, {'id': 'AL', 'value': 0.0}, {'id': 'AM', 'value': 0.0},
                {'id': 'AV', 'value': 0.0}, {'id': 'AU', 'value': 0.0}, {'id': 'JK', 'value': 0.0},
                {'id': 'AC', 'value': 0.0}, {'id': 'AT', 'value': ''}, {'id': 'AW', 'value': ''},
                {'id': 'BG', 'value': ''}, {'id': 'BM', 'value': ''}, {'id': 'DC', 'value': ''},
                {'id': 'DA', 'value': 'S'}, {'id': 'FC', 'value': ''}, {'id': 'FA', 'value': 'S'}
            ],
            '2031BIS': [
                {'id': 'EC', 'value': 0.0}, {'id': 'BF', 'value': 0.0}, {'id': 'BG', 'value': 0.0},
                {'id': 'GA', 'value': ''}, {'id': 'HA', 'value': 0.0}, {'id': 'HB', 'value': 0.0},
                {'id': 'HC', 'value': 0.0}, {'id': 'HD', 'value': 0.0}, {'id': 'AE', 'value': 0.0},
                {'id': 'AF', 'value': 0.0}, {'id': 'AD', 'value': 0.0}, {'id': 'AL', 'value': 0.0},
                {'id': 'AB', 'value': 0.0}, {'id': 'AC', 'value': 0.0}, {'id': 'SA', 'value': 0.0},
                {'id': 'TA', 'value': 0.0}, {'id': 'SB', 'value': 0.0}, {'id': 'TB', 'value': 0.0},
                {'id': 'SC', 'value': 0}, {'id': 'TC', 'value': 0}
            ],
            '2033A': [
                {'id': 'AR', 'value': -7500.0}, {'id': 'BR', 'value': 0}, {'id': 'CR', 'value': -7500.0},
                {'id': 'AE', 'value': -7500.0}, {'id': 'BE', 'value': 0}, {'id': 'CE', 'value': -7500.0},
                {'id': 'AA', 'value': 0}, {'id': 'BA', 'value': 0}, {'id': 'CA', 'value': 0},
                {'id': 'AB', 'value': 0}, {'id': 'BB', 'value': 0}, {'id': 'CB', 'value': 0},
                {'id': 'AC', 'value': -7500.0}, {'id': 'BC', 'value': 0}, {'id': 'CC', 'value': -7500.0},
                {'id': 'AD', 'value': 0}, {'id': 'BD', 'value': 0}, {'id': 'CD', 'value': 0},
                {'id': 'AQ', 'value': 0}, {'id': 'BQ', 'value': 0}, {'id': 'CQ', 'value': 0},
                {'id': 'AF', 'value': 0}, {'id': 'BF', 'value': 0}, {'id': 'CF', 'value': 0},
                {'id': 'AG', 'value': 0}, {'id': 'BG', 'value': 0}, {'id': 'CG', 'value': 0},
                {'id': 'AH', 'value': 0}, {'id': 'BH', 'value': None}, {'id': 'CH', 'value': 0},
                {'id': 'AJ', 'value': 0}, {'id': 'BJ', 'value': 0}, {'id': 'CJ', 'value': 0},
                {'id': 'AK', 'value': 0}, {'id': 'BK', 'value': 0}, {'id': 'CK', 'value': 0},
                {'id': 'AL', 'value': 0}, {'id': 'BL', 'value': 0}, {'id': 'CL', 'value': 0},
                {'id': 'AM', 'value': 0}, {'id': 'BM', 'value': None}, {'id': 'CM', 'value': 0},
                {'id': 'AP', 'value': 0}, {'id': 'BP', 'value': None}, {'id': 'CP', 'value': 0},
                {'id': 'FS', 'value': -7500.0}, {'id': 'FJ', 'value': -5000.0}, {'id': 'FA', 'value': 0},
                {'id': 'FB', 'value': 0}, {'id': 'FC', 'value': 0}, {'id': 'FD', 'value': 0},
                {'id': 'FE', 'value': 0}, {'id': 'FF', 'value': -5000.0}, {'id': 'FG', 'value': 0},
                {'id': 'AS', 'value': 0}, {'id': 'FH', 'value': 0}, {'id': 'FK', 'value': 0},
                {'id': 'FR', 'value': -2500.0}, {'id': 'FL', 'value': 0}, {'id': 'FM', 'value': 0},
                {'id': 'FN', 'value': 0}, {'id': 'AU', 'value': -2500.0}, {'id': 'EE', 'value': 0},
                {'id': 'FP', 'value': 0}, {'id': 'FQ', 'value': 0}, {'id': 'JB', 'value': 0},
                {'id': 'JC', 'value': 0}, {'id': 'HA', 'value': 0.0}, {'id': 'HB', 'value': 0.0},
                {'id': 'HC', 'value': 0}, {'id': 'EB', 'value': 0.0}, {'id': 'JA', 'value': 0.0},
                {'id': 'AT', 'value': 0}
            ],
            '2033B': [
                {'id': 'BW', 'value': 0}, {'id': 'BH', 'value': 0}, {'id': 'BA', 'value': 0},
                {'id': 'AA', 'value': 0}, {'id': 'BB', 'value': 0}, {'id': 'AB', 'value': 0},
                {'id': 'BC', 'value': 0}, {'id': 'AC', 'value': 0}, {'id': 'BD', 'value': 0},
                {'id': 'BE', 'value': 0}, {'id': 'BF', 'value': 0}, {'id': 'BG', 'value': 0},
                {'id': 'BV', 'value': 0}, {'id': 'BJ', 'value': 0}, {'id': 'BK', 'value': 0},
                {'id': 'BL', 'value': 0}, {'id': 'BM', 'value': 0}, {'id': 'BN', 'value': 0},
                {'id': 'GN', 'value': 0.0}, {'id': 'AN', 'value': 0.0}, {'id': 'BP', 'value': 0},
                {'id': 'AP', 'value': 0.0}, {'id': 'BQ', 'value': 0}, {'id': 'BR', 'value': 0},
                {'id': 'EE', 'value': 0.0}, {'id': 'BS', 'value': 0}, {'id': 'CK', 'value': 0.0},
                {'id': 'BT', 'value': 0}, {'id': 'BU', 'value': 0}, {'id': 'AU', 'value': 0.0},
                {'id': 'AV', 'value': 0.0}, {'id': 'BX', 'value': 0}, {'id': 'BY', 'value': 0},
                {'id': 'BZ', 'value': 0}, {'id': 'CA', 'value': 0}, {'id': 'AK', 'value': 0.0},
                {'id': 'AL', 'value': 0.0}, {'id': 'CB', 'value': 0}, {'id': 'CC', 'value': 0},
                {'id': 'CD', 'value': 0.0}, {'id': 'ED', 'value': 0.0}, {'id': 'CE', 'value': 0.0},
                {'id': 'CF', 'value': 0.0}, {'id': 'CG', 'value': 0.0}, {'id': 'CH', 'value': 0.0},
                {'id': 'CJ', 'value': 0.0}, {'id': 'FK', 'value': 0.0}, {'id': 'FJ', 'value': 0.0},
                {'id': 'AE', 'value': 0.0}, {'id': 'AD', 'value': 0.0}, {'id': 'AF', 'value': 0.0},
                {'id': 'AJ', 'value': 0.0}, {'id': 'AG', 'value': 0.0}, {'id': 'EK', 'value': 0.0},
                {'id': 'MC', 'value': 0.0}, {'id': 'MD', 'value': 0.0}, {'id': 'MF', 'value': 0.0},
                {'id': 'FA', 'value': 0.0}, {'id': 'FB', 'value': 0.0}, {'id': 'EU', 'value': 0.0},
                {'id': 'AQ', 'value': 0.0}, {'id': 'HL', 'value': 0.0}, {'id': 'MH', 'value': 0.0},
                {'id': 'AR', 'value': 0.0}, {'id': 'AY', 'value': 0.0}, {'id': 'EL', 'value': 0.0},
                {'id': 'FL', 'value': 0.0}, {'id': 'AH', 'value': 0.0}, {'id': 'AS', 'value': 0.0},
                {'id': 'AX', 'value': 0.0}, {'id': 'CX', 'value': 0.0}, {'id': 'CY', 'value': 0.0},
                {'id': 'CZ', 'value': 0.0}, {'id': 'AT', 'value': 0.0}, {'id': 'AZ', 'value': 0.0},
                {'id': 'CM', 'value': 0.0}, {'id': 'EM', 'value': 0.0}, {'id': 'CN', 'value': 0.0},
                {'id': 'EP', 'value': 0.0}, {'id': 'CR', 'value': 0.0}, {'id': 'ER', 'value': 0.0}
            ],
            '2033C': [
                {'id': 'AA', 'value': 0.0}, {'id': 'BA', 'value': 0.0}, {'id': 'CA', 'value': 0.0},
                {'id': 'DA', 'value': 0}, {'id': 'AB', 'value': 0.0}, {'id': 'BB', 'value': 0.0},
                {'id': 'CB', 'value': 0.0}, {'id': 'DB', 'value': 0}, {'id': 'AC', 'value': 0.0},
                {'id': 'BC', 'value': 0.0}, {'id': 'CC', 'value': 0.0}, {'id': 'DC', 'value': 0},
                {'id': 'AD', 'value': 0.0}, {'id': 'CD', 'value': 0.0}, {'id': 'BD', 'value': 0.0},
                {'id': 'DD', 'value': 0}, {'id': 'AE', 'value': 0.0}, {'id': 'BE', 'value': 0.0},
                {'id': 'CE', 'value': 0.0}, {'id': 'DE', 'value': 0}, {'id': 'AF', 'value': 0.0},
                {'id': 'BF', 'value': 0.0}, {'id': 'CF', 'value': 0.0}, {'id': 'DF', 'value': 0},
                {'id': 'AG', 'value': 0.0}, {'id': 'BG', 'value': 0.0}, {'id': 'CG', 'value': 0.0},
                {'id': 'DG', 'value': 0}, {'id': 'AH', 'value': 0.0}, {'id': 'BH', 'value': 0.0},
                {'id': 'CH', 'value': 0.0}, {'id': 'DH', 'value': 0}, {'id': 'AJ', 'value': 0.0},
                {'id': 'BJ', 'value': 0.0}, {'id': 'CJ', 'value': 0.0}, {'id': 'DJ', 'value': 0},
                {'id': 'AK', 'value': 0}, {'id': 'BK', 'value': 0}, {'id': 'CK', 'value': 0},
                {'id': 'DK', 'value': 0}, {'id': 'AL', 'value': 0.0}, {'id': 'AM', 'value': 0.0},
                {'id': 'AN', 'value': 0.0}, {'id': 'AP', 'value': 0}, {'id': 'FA', 'value': 0.0},
                {'id': 'GA', 'value': 0.0}, {'id': 'HA', 'value': 0.0}, {'id': 'JA', 'value': 0},
                {'id': 'FB', 'value': 0.0}, {'id': 'GB', 'value': 0.0}, {'id': 'HB', 'value': 0.0},
                {'id': 'JB', 'value': 0}, {'id': 'FC', 'value': 0.0}, {'id': 'GC', 'value': 0.0},
                {'id': 'HC', 'value': 0.0}, {'id': 'JC', 'value': 0}, {'id': 'FD', 'value': 0.0},
                {'id': 'GD', 'value': 0.0}, {'id': 'HD', 'value': 0.0}, {'id': 'JD', 'value': 0},
                {'id': 'FE', 'value': 0.0}, {'id': 'GE', 'value': 0.0}, {'id': 'HE', 'value': 0.0},
                {'id': 'JE', 'value': 0}, {'id': 'FF', 'value': 0.0}, {'id': 'GF', 'value': 0.0},
                {'id': 'HF', 'value': 0.0}, {'id': 'JF', 'value': 0}, {'id': 'FG', 'value': 0.0},
                {'id': 'GG', 'value': 0.0}, {'id': 'HG', 'value': 0.0}, {'id': 'JG', 'value': 0},
                {'id': 'FH', 'value': 0}, {'id': 'GH', 'value': 0}, {'id': 'HH', 'value': 0},
                {'id': 'JH', 'value': 0}, {'id': 'LL', 'value': 0}, {'id': 'ML', 'value': 0},
                {'id': 'NL', 'value': 0}, {'id': 'PL', 'value': 0}, {'id': 'QL', 'value': 0},
                {'id': 'SL', 'value': 0}, {'id': 'UL', 'value': 0}, {'id': 'TL', 'value': 0},
                {'id': 'QP', 'value': 0.0}, {'id': 'QM', 'value': 0.0}, {'id': 'SM', 'value': 0.0},
                {'id': 'UM', 'value': 0.0}, {'id': 'TM', 'value': 0.0}, {'id': 'QN', 'value': 0},
                {'id': 'SN', 'value': 0}, {'id': 'UN', 'value': 0}, {'id': 'TN', 'value': 0}
            ],
            '2033D': [
                {'id': 'AA', 'value': 0.0}, {'id': 'BA', 'value': 0.0}, {'id': 'DA', 'value': 0.0},
                {'id': 'EA', 'value': 0.0}, {'id': 'AN', 'value': 0.0}, {'id': 'BN', 'value': 0.0},
                {'id': 'CN', 'value': 0.0}, {'id': 'DN', 'value': 0.0}, {'id': 'AB', 'value': 0.0},
                {'id': 'BB', 'value': 0.0}, {'id': 'CB', 'value': 0.0}, {'id': 'DB', 'value': 0.0},
                {'id': 'AC', 'value': 0.0}, {'id': 'BC', 'value': 0.0}, {'id': 'CC', 'value': 0.0},
                {'id': 'DC', 'value': 0.0}, {'id': 'AD', 'value': 0.0}, {'id': 'BD', 'value': 0.0},
                {'id': 'CD', 'value': 0.0}, {'id': 'DD', 'value': 0.0}, {'id': 'AE', 'value': 0.0},
                {'id': 'BE', 'value': 0.0}, {'id': 'CE', 'value': 0.0}, {'id': 'DE', 'value': 0.0},
                {'id': 'AF', 'value': 0.0}, {'id': 'BF', 'value': 0.0}, {'id': 'CF', 'value': 0.0},
                {'id': 'DF', 'value': 0.0}, {'id': 'AG', 'value': 0.0}, {'id': 'BG', 'value': 0.0},
                {'id': 'CG', 'value': 0.0}, {'id': 'DG', 'value': 0.0}, {'id': 'AH', 'value': 0},
                {'id': 'BH', 'value': 0}, {'id': 'CH', 'value': 0}, {'id': 'DH', 'value': 0},
                {'id': 'AY', 'value': 0.0}, {'id': 'AZ', 'value': 0.0}, {'id': 'FA', 'value': 0.0},
                {'id': 'EB', 'value': 0.0}, {'id': 'FB', 'value': 0.0}, {'id': 'EC', 'value': 0.0},
                {'id': 'FC', 'value': 0.0}, {'id': 'ED', 'value': 0.0}, {'id': 'FD', 'value': 0.0},
                {'id': 'EE', 'value': 0.0}, {'id': 'FE', 'value': 0.0}, {'id': 'EF', 'value': 0.0},
                {'id': 'FF', 'value': 0.0}, {'id': 'EG', 'value': 0.0}, {'id': 'FG', 'value': 0.0},
                {'id': 'EH', 'value': 0}, {'id': 'FH', 'value': 0},
                {'id': 'HJ', 'occurences': [{'value': None, 'occurence_number': '0001'}]},
                {'id': 'HA', 'value': 0.0}, {'id': 'HH', 'value': 0}, {'id': 'PG', 'value': 0.0},
                {'id': 'AL', 'value': 0.0}, {'id': 'AM', 'value': 0.0}, {'id': 'PH', 'value': 0.0},
                {'id': 'PJ', 'value': 0}, {'id': 'MG', 'value': 0.0}, {'id': 'MH', 'value': 0},
                {'id': 'AP', 'value': 0.0}, {'id': 'AW', 'value': 0.0}, {'id': 'AX', 'value': 0.0},
                {'id': 'AS', 'value': 0.0}, {'id': 'AT', 'value': 0.0}, {'id': 'AU', 'value': 0.0},
                {'id': 'AJ', 'value': 0.0}, {'id': 'AK', 'value': 0.0}, {'id': 'DP', 'value': 0.0},
                {'id': 'DQ', 'value': 0.0}, {'id': 'DR', 'value': 0}, {'id': 'DS', 'value': 0}
            ],
            '2033E': [
                {'id': 'AA', 'value': 0.0}, {'id': 'AB', 'value': 0.0}, {'id': 'AC', 'value': 0.0},
                {'id': 'AD', 'value': 0.0}, {'id': 'GA', 'value': 0}, {'id': 'AF', 'value': 0},
                {'id': 'AJ', 'value': 0.0}, {'id': 'EN', 'value': 0.0}, {'id': 'EP', 'value': 0.0},
                {'id': 'AN', 'value': 0}, {'id': 'EL', 'value': 0.0}, {'id': 'EE', 'value': 0.0},
                {'id': 'EF', 'value': 0.0}, {'id': 'AK', 'value': 0.0}, {'id': 'GD', 'value': 0.0},
                {'id': 'AP', 'value': 0}, {'id': 'AL', 'value': 0.0}, {'id': 'AM', 'value': 0.0},
                {'id': 'FE', 'value': 0.0}, {'id': 'FJ', 'value': 0.0}, {'id': 'FH', 'value': 0.0},
                {'id': 'FK', 'value': 0.0}, {'id': 'GC', 'value': 0.0}, {'id': 'FN', 'value': 0.0},
                {'id': 'FM', 'value': 0.0}, {'id': 'GB', 'value': 0.0}, {'id': 'KA', 'value': ''},
                {'id': 'KD', 'value': 0.0}, {'id': 'AH', 'value': 0.0}, {'id': 'AQ', 'value': 0.0},
                {'id': 'KB', 'value': ''}, {'id': 'KC', 'value': ''}, {'id': 'KG', 'value': ''}
            ],
            '2033F': [
                {'id': 'GX', 'value': 0}, {'id': 'GT', 'value': 0.0}, {'id': 'GV', 'value': 0.0},
                {'id': 'GY', 'value': 0}, {'id': 'GU', 'value': 0.0}, {'id': 'GW', 'value': 0.0}
            ],
            '2033G': [
                {'id': 'GT', 'value': 0}
            ],
            '2065': [
                {'id': 'PH', 'value': ''}, {'id': 'PE', 'value': ''}, {'id': 'AS', 'value': ''},
                {'id': 'AC', 'value': ''}, {'id': 'AK', 'value': ''}, {'id': 'PA', 'value': ''},
                    {'id': 'PD', 'value': None}, {'id': 'AQ', 'ftx_1': ''}, {'id': 'AP', 'value': ''},
                {'id': 'HA', 'value': 0.0}, {'id': 'AT', 'value': 0.0}, {'id': 'LC', 'value': 0.0},
                {'id': 'AL', 'value': 0.0}, {'id': 'LT', 'value': 0.0}, {'id': 'LN', 'value': 0.0},
                {'id': 'AU', 'value': 0.0}, {'id': 'LW', 'value': 0.0}, {'id': 'LX', 'value': 0.0},
                {'id': 'AX', 'value': ''}, {'id': 'LQ', 'value': ''}, {'id': 'MA', 'value': ''},
                {'id': 'BY', 'value': ''}, {'id': 'LY', 'value': ''}, {'id': 'AF', 'value': ''},
                {'id': 'HD', 'value': ''}, {'id': 'AW', 'value': ''}, {'id': 'LL', 'value': ''},
                {'id': 'AM', 'value': ''}, {'id': 'BA', 'value': ''}, {'id': 'HJ', 'value': ''},
                {'id': 'BB', 'value': 0.0}, {'id': 'LV', 'value': 0.0}, {'id': 'PF', 'value': ''},
                {'id': 'BE', 'value': 0.0}, {'id': 'BF', 'value': 0.0}, {'id': 'JA', 'value': 0.0},
                {'id': 'AR', 'value': ''}, {'id': 'CA', 'value': ''}, {'id': 'BN', 'value': 'S'},
                {'id': 'EA', 'value': ''}, {'id': 'BQ', 'value': 'S'}
            ],
            '2065BIS': [
                {'id': 'AA', 'value': 0.0}, {'id': 'AB', 'value': 0.0}, {'id': 'AC', 'value': 0.0},
                {'id': 'AD', 'value': 0.0},
                {'id': 'BE', 'occurences': [
                    {'ftx_1': '', 'occurence_number': '0001'}, {'ftx_1': '', 'occurence_number': '0002'},
                    {'ftx_1': '', 'occurence_number': '0003'}, {'ftx_1': '', 'occurence_number': '0004'}
                ]},
                {'id': 'AE', 'occurences': [
                    {'value': 0.0, 'occurence_number': '0001'}, {'value': 0.0, 'occurence_number': '0002'},
                    {'value': 0.0, 'occurence_number': '0003'}, {'value': 0.0, 'occurence_number': '0004'}
                ]},
                {'id': 'AK', 'value': 0.0}, {'id': 'AL', 'value': 0.0}, {'id': 'AJ', 'value': 0},
                {'id': 'CV', 'value': ''}, {'id': 'CX', 'value': 0.0}, {'id': 'CY', 'value': 0.0},
                {'id': 'AH', 'value': 0.0}, {'id': 'CZ', 'value': 0.0}, {'id': 'AQ', 'value': 0.0},
                {'id': 'AM', 'value': 0.0}, {'id': 'DA', 'value': 0.0}, {'id': 'AR', 'value': 0.0},
                {'id': 'AN', 'value': 0.0}, {'id': 'DB', 'value': 0.0}, {'id': 'AS', 'value': 0.0},
                {'id': 'AP', 'value': 0.0}, {'id': 'DC', 'value': 0.0}, {'id': 'AT', 'value': 0.0},
                {'id': 'AF', 'value': 0.0}, {'id': 'AG', 'value': 0.0}
            ],
            '2069RCI': [
                {'id': 'AE', 'value': ''}, {'id': 'AD', 'value': ''}, {'id': 'BE', 'value': ''},
                {'id': 'BB', 'occurences': [
                    {'value': 0.0, 'occurence_number': '0001'}, {'value': 0.0, 'occurence_number': '0002'},
                    {'value': 0.0, 'occurence_number': '0003'}, {'value': 0.0, 'occurence_number': '0004'},
                    {'value': 0.0, 'occurence_number': '0005'}, {'value': 0.0, 'occurence_number': '0006'},
                    {'value': 0.0, 'occurence_number': '0007'}, {'value': 0.0, 'occurence_number': '0008'},
                    {'value': 0.0, 'occurence_number': '0009'}, {'value': 0.0, 'occurence_number': '0010'},
                    {'value': 0.0, 'occurence_number': '0011'}, {'value': 0.0, 'occurence_number': '0012'},
                    {'value': 0.0, 'occurence_number': '0013'}, {'value': 0.0, 'occurence_number': '0014'},
                    {'value': 0.0, 'occurence_number': '0015'}, {'value': 0.0, 'occurence_number': '0016'},
                    {'value': 0.0, 'occurence_number': '0017'}, {'value': 0.0, 'occurence_number': '0018'},
                    {'value': 0.0, 'occurence_number': '0019'}
                ]},
                {'id': 'BA', 'occurences': [
                    {'value': 'VEL', 'occurence_number': '0001'}, {'value': 'AUT', 'occurence_number': '0002'},
                    {'value': 'AUT', 'occurence_number': '0003'}, {'value': 'MEC', 'occurence_number': '0004'},
                    {'value': 'CIC', 'occurence_number': '0005'}, {'value': 'COM', 'occurence_number': '0006'},
                    {'value': '2LI', 'occurence_number': '0007'}, {'value': 'AUT', 'occurence_number': '0008'},
                    {'value': 'RAC', 'occurence_number': '0009'}, {'value': 'CIN', 'occurence_number': '0010'},
                    {'value': 'AUD', 'occurence_number': '0011'}, {'value': 'CCI', 'occurence_number': '0012'},
                    {'value': 'CSV', 'occurence_number': '0013'}, {'value': 'RTD', 'occurence_number': '0014'},
                    {'value': 'REB', 'occurence_number': '0015'}, {'value': 'AUT', 'occurence_number': '0016'},
                    {'value': 'HVE', 'occurence_number': '0017'}, {'value': '3IV', 'occurence_number': '0018'},
                    {'value': 'PAM', 'occurence_number': '0019'}
                ]},
                {'id': 'AU', 'value': 0.0}, {'id': 'BP', 'value': 0.0}, {'id': 'AV', 'value': 0.0},
                {'id': 'AM', 'value': 0.0}, {'id': 'AN', 'value': 0.0}, {'id': 'AP', 'value': 0.0},
                {'id': 'BD', 'occurences': [
                    {'value': 0.0, 'occurence_number': '0001'}, {'value': 0.0, 'occurence_number': '0002'},
                    {'value': 0.0, 'occurence_number': '0003'}, {'value': 0.0, 'occurence_number': '0004'},
                    {'value': 0.0, 'occurence_number': '0005'}, {'value': 0.0, 'occurence_number': '0006'},
                    {'value': 0.0, 'occurence_number': '0007'}, {'value': 0.0, 'occurence_number': '0008'},
                    {'value': 0.0, 'occurence_number': '0009'}, {'value': 0.0, 'occurence_number': '0010'},
                    {'value': 0.0, 'occurence_number': '0011'}, {'value': 0.0, 'occurence_number': '0012'},
                    {'value': 0.0, 'occurence_number': '0013'}, {'value': 0.0, 'occurence_number': '0014'},
                    {'value': 0.0, 'occurence_number': '0015'}, {'value': 0.0, 'occurence_number': '0016'}
                ]},
                {'id': 'BC', 'occurences': [
                    {'value': 'PVM', 'occurence_number': '0001'}, {'value': 'COR', 'occurence_number': '0002'},
                    {'value': 'CIR', 'occurence_number': '0003'}, {'value': 'CRC', 'occurence_number': '0004'},
                    {'value': 'TZM', 'occurence_number': '0005'}, {'value': 'FAM', 'occurence_number': '0006'},
                    {'value': 'BIO', 'occurence_number': '0007'}, {'value': 'PHO', 'occurence_number': '0008'},
                    {'value': 'ART', 'occurence_number': '0009'}, {'value': 'CJV', 'occurence_number': '0010'},
                    {'value': 'EOM', 'occurence_number': '0011'}, {'value': 'CPE', 'occurence_number': '0012'},
                    {'value': 'PTR', 'occurence_number': '0013'}, {'value': 'RTA', 'occurence_number': '0014'},
                    {'value': 'CIO', 'occurence_number': '0015'}, {'value': 'COL', 'occurence_number': '0016'}
                ]},
                {'id': 'BK', 'value': 0.0}, {'id': 'AT', 'ftx_1': ''}
            ]
        }
        for report_name, config in REPORTS_TO_EXPORT.items():
            lines = []
            for report in config['reports']:
                section_id = self.env.ref(f'l10n_fr_reports.{report}').id
                section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
                section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
                lines.extend(section._get_lines(section_options))
            mapper = self.env[config['handler']]._get_code_to_edi_id()
            vals = self.env['l10n_fr_reports.fiscal_declaration_handler']._export_report_data(lines, mapper)

            with self.subTest(vals=vals):
                self.assertEqual(vals, expected[report_name])

    def test_data_extraction_dynamic_and_fixed_line(self):
        self._generate_move(amount=3000, move_type='out_invoice').action_post()
        self._generate_move(partner=self.partner_b, amount=2500).action_post()
        self._generate_move(partner=self.partner_b, amount=2000, move_type='out_invoice', account=self.copyright_account).action_post()
        self._generate_move(partner=self.partner_a, amount=5000, account=self.fees_account).action_post()
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        expected = {
            '2031': [
                {'id': 'AA', 'value': ''}, {'id': 'AD', 'value': ''}, {'id': 'KY', 'value': ''},
                {'id': 'AB', 'ftx_1': ''}, {'id': 'BT', 'value': ''}, {'id': 'CA', 'value': 0.0},
                {'id': 'CB', 'value': 0.0}, {'id': 'CC', 'value': 0.0}, {'id': 'CD', 'value': 0.0},
                {'id': 'CE', 'value': 0}, {'id': 'CF', 'value': 0.0}, {'id': 'CH', 'value': 0},
                {'id': 'CJ', 'value': 0}, {'id': 'CK', 'value': 0}, {'id': 'CL', 'value': 0},
                {'id': 'CM', 'value': 0}, {'id': 'AQ', 'value': 0.0}, {'id': 'AE', 'value': 0.0},
                {'id': 'AF', 'value': 0.0}, {'id': 'CP', 'value': 0.0}, {'id': 'KG', 'value': 0.0},
                {'id': 'CN', 'value': 0.0}, {'id': 'CV', 'value': 0.0}, {'id': 'KM', 'value': 0.0},
                {'id': 'AP', 'value': 0.0}, {'id': 'JA', 'value': ''}, {'id': 'KX', 'value': ''},
                {'id': 'KQ', 'value': ''}, {'id': 'JL', 'value': ''}, {'id': 'JB', 'value': ''},
                {'id': 'KR', 'value': ''}, {'id': 'KK', 'value': ''}, {'id': 'AG', 'value': ''},
                {'id': 'AR', 'value': ''}, {'id': 'AY', 'value': ''}, {'id': 'CQ', 'value': 0.0},
                {'id': 'CR', 'value': 0.0}, {'id': 'KZ', 'value': ''}, {'id': 'CW', 'value': 0},
                {'id': 'CX', 'value': 0}, {'id': 'AL', 'value': 0.0}, {'id': 'AM', 'value': 0.0},
                {'id': 'AV', 'value': 0.0}, {'id': 'AU', 'value': 0.0}, {'id': 'JK', 'value': 0.0},
                {'id': 'AC', 'value': 0.0}, {'id': 'AT', 'value': ''}, {'id': 'AW', 'value': ''},
                {'id': 'BG', 'value': ''}, {'id': 'BM', 'value': ''}, {'id': 'DC', 'value': ''},
                {'id': 'DA', 'value': 'S'}, {'id': 'FC', 'value': ''}, {'id': 'FA', 'value': 'S'}
            ],
            '2031BIS': [
                {'id': 'EC', 'value': 0.0}, {'id': 'BF', 'value': 0.0}, {'id': 'BG', 'value': 0.0},
                {'id': 'GA', 'value': ''}, {'id': 'HA', 'value': 0.0}, {'id': 'HB', 'value': 0.0},
                {'id': 'HC', 'value': 0.0}, {'id': 'HD', 'value': 0.0}, {'id': 'AE', 'value': 0.0},
                {'id': 'AF', 'value': 0.0}, {'id': 'AD', 'value': 0.0}, {'id': 'AL', 'value': 0.0},
                {'id': 'AB', 'value': 0.0}, {'id': 'AC', 'value': 0.0}, {'id': 'SA', 'value': 0.0},
                {'id': 'TA', 'value': 0.0}, {'id': 'SB', 'value': 0.0}, {'id': 'TB', 'value': 0.0},
                {'id': 'SC', 'value': 0}, {'id': 'TC', 'value': 0},
                {'id': 'AW', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'AV', 'occurences': [{'address': {'postal_code': '', 'city': '', 'country_code': None}, 'occurence_number': '0001'}]},
                {'id': 'AU', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'AX', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'AP', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'AQ', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'AA', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'AH', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'AJ', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'FA', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'KA', 'ftx_1': ''},
                {'id': 'KB', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'KC', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'KD', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]}
            ],
            '2033A': [
                {'id': 'AR', 'value': -7500.0}, {'id': 'BR', 'value': 0}, {'id': 'CR', 'value': -7500.0},
                {'id': 'AE', 'value': -7500.0}, {'id': 'BE', 'value': 0}, {'id': 'CE', 'value': -7500.0},
                {'id': 'AA', 'value': 0}, {'id': 'BA', 'value': 0}, {'id': 'CA', 'value': 0},
                {'id': 'AB', 'value': 0}, {'id': 'BB', 'value': 0}, {'id': 'CB', 'value': 0},
                {'id': 'AC', 'value': -7500.0}, {'id': 'BC', 'value': 0}, {'id': 'CC', 'value': -7500.0},
                {'id': 'AD', 'value': 0}, {'id': 'BD', 'value': 0}, {'id': 'CD', 'value': 0},
                {'id': 'AQ', 'value': 0}, {'id': 'BQ', 'value': 0}, {'id': 'CQ', 'value': 0},
                {'id': 'AF', 'value': 0}, {'id': 'BF', 'value': 0}, {'id': 'CF', 'value': 0},
                {'id': 'AG', 'value': 0}, {'id': 'BG', 'value': 0}, {'id': 'CG', 'value': 0},
                {'id': 'AH', 'value': 0}, {'id': 'BH', 'value': None}, {'id': 'CH', 'value': 0},
                {'id': 'AJ', 'value': 0}, {'id': 'BJ', 'value': 0}, {'id': 'CJ', 'value': 0},
                {'id': 'AK', 'value': 0}, {'id': 'BK', 'value': 0}, {'id': 'CK', 'value': 0},
                {'id': 'AL', 'value': 0}, {'id': 'BL', 'value': 0}, {'id': 'CL', 'value': 0},
                {'id': 'AM', 'value': 0}, {'id': 'BM', 'value': None}, {'id': 'CM', 'value': 0},
                {'id': 'AP', 'value': 0}, {'id': 'BP', 'value': None}, {'id': 'CP', 'value': 0},
                {'id': 'FS', 'value': -7500.0}, {'id': 'FJ', 'value': -5000.0}, {'id': 'FA', 'value': 0},
                {'id': 'FB', 'value': 0}, {'id': 'FC', 'value': 0}, {'id': 'FD', 'value': 0},
                {'id': 'FE', 'value': 0}, {'id': 'FF', 'value': -5000.0}, {'id': 'FG', 'value': 0},
                {'id': 'AS', 'value': 0}, {'id': 'FH', 'value': 0}, {'id': 'FK', 'value': 0},
                {'id': 'FR', 'value': -2500.0}, {'id': 'FL', 'value': 0}, {'id': 'FM', 'value': 0},
                {'id': 'FN', 'value': 0}, {'id': 'AU', 'value': -2500.0}, {'id': 'EE', 'value': 0},
                {'id': 'FP', 'value': 0}, {'id': 'FQ', 'value': 0}, {'id': 'JB', 'value': 0},
                {'id': 'JC', 'value': 0}, {'id': 'HA', 'value': 0.0}, {'id': 'HB', 'value': 0.0},
                {'id': 'HC', 'value': 0}, {'id': 'EB', 'value': 0.0}, {'id': 'JA', 'value': 0.0},
                {'id': 'AT', 'value': 0}
            ],
            '2033B': [
                {'id': 'BW', 'value': 0}, {'id': 'BH', 'value': 0}, {'id': 'BA', 'value': 0},
                {'id': 'AA', 'value': 0}, {'id': 'BB', 'value': 0}, {'id': 'AB', 'value': 0},
                {'id': 'BC', 'value': 0}, {'id': 'AC', 'value': 0}, {'id': 'BD', 'value': 0},
                {'id': 'BE', 'value': 0}, {'id': 'BF', 'value': 0}, {'id': 'BG', 'value': 0},
                {'id': 'BV', 'value': 0}, {'id': 'BJ', 'value': 0}, {'id': 'BK', 'value': 0},
                {'id': 'BL', 'value': 0}, {'id': 'BM', 'value': 0}, {'id': 'BN', 'value': 0},
                {'id': 'GN', 'value': 0.0}, {'id': 'AN', 'value': 0.0}, {'id': 'BP', 'value': 0},
                {'id': 'AP', 'value': 0.0}, {'id': 'BQ', 'value': 0}, {'id': 'BR', 'value': 0},
                {'id': 'EE', 'value': 0.0}, {'id': 'BS', 'value': 0}, {'id': 'CK', 'value': 0.0},
                {'id': 'BT', 'value': 0}, {'id': 'BU', 'value': 0}, {'id': 'AU', 'value': 0.0},
                {'id': 'AV', 'value': 0.0}, {'id': 'BX', 'value': 0}, {'id': 'BY', 'value': 0},
                {'id': 'BZ', 'value': 0}, {'id': 'CA', 'value': 0}, {'id': 'AK', 'value': 0.0},
                {'id': 'AL', 'value': 0.0}, {'id': 'CB', 'value': 0}, {'id': 'CC', 'value': 0},
                {'id': 'CD', 'value': 0.0}, {'id': 'ED', 'value': 0.0}, {'id': 'CE', 'value': 0.0},
                {'id': 'CF', 'value': 0.0}, {'id': 'CG', 'value': 0.0}, {'id': 'CH', 'value': 0.0},
                {'id': 'CJ', 'value': 0.0}, {'id': 'FK', 'value': 0.0}, {'id': 'FJ', 'value': 0.0},
                {'id': 'AE', 'value': 0.0}, {'id': 'AD', 'value': 0.0}, {'id': 'AF', 'value': 0.0},
                {'id': 'AJ', 'value': 0.0}, {'id': 'AG', 'value': 0.0}, {'id': 'EK', 'value': 0.0},
                {'id': 'MC', 'value': 0.0}, {'id': 'MD', 'value': 0.0}, {'id': 'MF', 'value': 0.0},
                {'id': 'FA', 'value': 0.0}, {'id': 'FB', 'value': 0.0}, {'id': 'EU', 'value': 0.0},
                {'id': 'AQ', 'value': 0.0}, {'id': 'HL', 'value': 0.0}, {'id': 'MH', 'value': 0.0},
                {'id': 'AR', 'value': 0.0}, {'id': 'AY', 'value': 0.0}, {'id': 'EL', 'value': 0.0},
                {'id': 'FL', 'value': 0.0}, {'id': 'AH', 'value': 0.0}, {'id': 'AS', 'value': 0.0},
                {'id': 'AX', 'value': 0.0}, {'id': 'CX', 'value': 0.0}, {'id': 'CY', 'value': 0.0},
                {'id': 'CZ', 'value': 0.0}, {'id': 'AT', 'value': 0.0}, {'id': 'AZ', 'value': 0.0},
                {'id': 'CM', 'value': 0.0}, {'id': 'EM', 'value': 0.0}, {'id': 'CN', 'value': 0.0},
                {'id': 'EP', 'value': 0.0}, {'id': 'CR', 'value': 0.0}, {'id': 'ER', 'value': 0.0}
            ],
            '2033C': [
                {'id': 'AA', 'value': 0.0}, {'id': 'BA', 'value': 0.0}, {'id': 'CA', 'value': 0.0},
                {'id': 'DA', 'value': 0}, {'id': 'AB', 'value': 0.0}, {'id': 'BB', 'value': 0.0},
                {'id': 'CB', 'value': 0.0}, {'id': 'DB', 'value': 0}, {'id': 'AC', 'value': 0.0},
                {'id': 'BC', 'value': 0.0}, {'id': 'CC', 'value': 0.0}, {'id': 'DC', 'value': 0},
                {'id': 'AD', 'value': 0.0}, {'id': 'CD', 'value': 0.0}, {'id': 'BD', 'value': 0.0},
                {'id': 'DD', 'value': 0}, {'id': 'AE', 'value': 0.0}, {'id': 'BE', 'value': 0.0},
                {'id': 'CE', 'value': 0.0}, {'id': 'DE', 'value': 0}, {'id': 'AF', 'value': 0.0},
                {'id': 'BF', 'value': 0.0}, {'id': 'CF', 'value': 0.0}, {'id': 'DF', 'value': 0},
                {'id': 'AG', 'value': 0.0}, {'id': 'BG', 'value': 0.0}, {'id': 'CG', 'value': 0.0},
                {'id': 'DG', 'value': 0}, {'id': 'AH', 'value': 0.0}, {'id': 'BH', 'value': 0.0},
                {'id': 'CH', 'value': 0.0}, {'id': 'DH', 'value': 0}, {'id': 'AJ', 'value': 0.0},
                {'id': 'BJ', 'value': 0.0}, {'id': 'CJ', 'value': 0.0}, {'id': 'DJ', 'value': 0},
                {'id': 'AK', 'value': 0}, {'id': 'BK', 'value': 0}, {'id': 'CK', 'value': 0},
                {'id': 'DK', 'value': 0}, {'id': 'AL', 'value': 0.0}, {'id': 'AM', 'value': 0.0},
                {'id': 'AN', 'value': 0.0}, {'id': 'AP', 'value': 0}, {'id': 'FA', 'value': 0.0},
                {'id': 'GA', 'value': 0.0}, {'id': 'HA', 'value': 0.0}, {'id': 'JA', 'value': 0},
                {'id': 'FB', 'value': 0.0}, {'id': 'GB', 'value': 0.0}, {'id': 'HB', 'value': 0.0},
                {'id': 'JB', 'value': 0}, {'id': 'FC', 'value': 0.0}, {'id': 'GC', 'value': 0.0},
                {'id': 'HC', 'value': 0.0}, {'id': 'JC', 'value': 0}, {'id': 'FD', 'value': 0.0},
                {'id': 'GD', 'value': 0.0}, {'id': 'HD', 'value': 0.0}, {'id': 'JD', 'value': 0},
                {'id': 'FE', 'value': 0.0}, {'id': 'GE', 'value': 0.0}, {'id': 'HE', 'value': 0.0},
                {'id': 'JE', 'value': 0}, {'id': 'FF', 'value': 0.0}, {'id': 'GF', 'value': 0.0},
                {'id': 'HF', 'value': 0.0}, {'id': 'JF', 'value': 0}, {'id': 'FG', 'value': 0.0},
                {'id': 'GG', 'value': 0.0}, {'id': 'HG', 'value': 0.0}, {'id': 'JG', 'value': 0},
                {'id': 'FH', 'value': 0}, {'id': 'GH', 'value': 0}, {'id': 'HH', 'value': 0},
                {'id': 'JH', 'value': 0}, {'id': 'LL', 'value': 0}, {'id': 'ML', 'value': 0},
                {'id': 'NL', 'value': 0}, {'id': 'PL', 'value': 0}, {'id': 'QL', 'value': 0},
                {'id': 'SL', 'value': 0}, {'id': 'UL', 'value': 0}, {'id': 'TL', 'value': 0},
                {'id': 'QP', 'value': 0.0}, {'id': 'QM', 'value': 0.0}, {'id': 'SM', 'value': 0.0},
                {'id': 'UM', 'value': 0.0}, {'id': 'TM', 'value': 0.0}, {'id': 'QN', 'value': 0},
                {'id': 'SN', 'value': 0}, {'id': 'UN', 'value': 0}, {'id': 'TN', 'value': 0},
                {'id': 'KA', 'occurences': [{'ftx_1': '', 'occurence_number': '0001'}]},
                {'id': 'LA', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'MA', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'NA', 'occurences': [{'value': 0, 'occurence_number': '0001'}]},
                {'id': 'PA', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'QA', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'SA', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'UA', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'TA', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]}
            ],
            '2033D': [
                {'id': 'AA', 'value': 0.0}, {'id': 'BA', 'value': 0.0}, {'id': 'DA', 'value': 0.0},
                {'id': 'EA', 'value': 0.0}, {'id': 'AN', 'value': 0.0}, {'id': 'BN', 'value': 0.0},
                {'id': 'CN', 'value': 0.0}, {'id': 'DN', 'value': 0.0}, {'id': 'AB', 'value': 0.0},
                {'id': 'BB', 'value': 0.0}, {'id': 'CB', 'value': 0.0}, {'id': 'DB', 'value': 0.0},
                {'id': 'AC', 'value': 0.0}, {'id': 'BC', 'value': 0.0}, {'id': 'CC', 'value': 0.0},
                {'id': 'DC', 'value': 0.0}, {'id': 'AD', 'value': 0.0}, {'id': 'BD', 'value': 0.0},
                {'id': 'CD', 'value': 0.0}, {'id': 'DD', 'value': 0.0}, {'id': 'AE', 'value': 0.0},
                {'id': 'BE', 'value': 0.0}, {'id': 'CE', 'value': 0.0}, {'id': 'DE', 'value': 0.0},
                {'id': 'AF', 'value': 0.0}, {'id': 'BF', 'value': 0.0}, {'id': 'CF', 'value': 0.0},
                {'id': 'DF', 'value': 0.0}, {'id': 'AG', 'value': 0.0}, {'id': 'BG', 'value': 0.0},
                {'id': 'CG', 'value': 0.0}, {'id': 'DG', 'value': 0.0}, {'id': 'AH', 'value': 0},
                {'id': 'BH', 'value': 0}, {'id': 'CH', 'value': 0}, {'id': 'DH', 'value': 0},
                {'id': 'AY', 'value': 0.0}, {'id': 'AZ', 'value': 0.0}, {'id': 'FA', 'value': 0.0},
                {'id': 'EB', 'value': 0.0}, {'id': 'FB', 'value': 0.0}, {'id': 'EC', 'value': 0.0},
                {'id': 'FC', 'value': 0.0}, {'id': 'ED', 'value': 0.0}, {'id': 'FD', 'value': 0.0},
                {'id': 'EE', 'value': 0.0}, {'id': 'FE', 'value': 0.0}, {'id': 'EF', 'value': 0.0},
                {'id': 'FF', 'value': 0.0}, {'id': 'EG', 'value': 0.0}, {'id': 'FG', 'value': 0.0},
                {'id': 'EH', 'value': 0}, {'id': 'FH', 'value': 0},
                {'id': 'HJ', 'occurences': [{'value': None, 'occurence_number': '0001'}, {'value': '', 'occurence_number': '0002'}]},
                {'id': 'HA', 'value': 0.0}, {'id': 'HH', 'value': 0}, {'id': 'PG', 'value': 0.0},
                {'id': 'AL', 'value': 0.0}, {'id': 'AM', 'value': 0.0}, {'id': 'PH', 'value': 0.0},
                {'id': 'PJ', 'value': 0}, {'id': 'MG', 'value': 0.0}, {'id': 'MH', 'value': 0},
                {'id': 'AP', 'value': 0.0}, {'id': 'AW', 'value': 0.0}, {'id': 'AX', 'value': 0.0},
                {'id': 'AS', 'value': 0.0}, {'id': 'AT', 'value': 0.0}, {'id': 'AU', 'value': 0.0},
                {'id': 'AJ', 'value': 0.0}, {'id': 'AK', 'value': 0.0}, {'id': 'DP', 'value': 0.0},
                {'id': 'DQ', 'value': 0.0}, {'id': 'DR', 'value': 0}, {'id': 'DS', 'value': 0},
                {'id': 'GJ', 'occurences': [{'ftx_1': 0.0, 'occurence_number': '0001'}]}
            ],
            '2033E': [
                {'id': 'AA', 'value': 0.0}, {'id': 'AB', 'value': 0.0}, {'id': 'AC', 'value': 0.0},
                {'id': 'AD', 'value': 0.0}, {'id': 'GA', 'value': 0}, {'id': 'AF', 'value': 0},
                {'id': 'AJ', 'value': 0.0}, {'id': 'EN', 'value': 0.0}, {'id': 'EP', 'value': 0.0},
                {'id': 'AN', 'value': 0}, {'id': 'EL', 'value': 0.0}, {'id': 'EE', 'value': 0.0},
                {'id': 'EF', 'value': 0.0}, {'id': 'AK', 'value': 0.0}, {'id': 'GD', 'value': 0.0},
                {'id': 'AP', 'value': 0}, {'id': 'AL', 'value': 0.0}, {'id': 'AM', 'value': 0.0},
                {'id': 'FE', 'value': 0.0}, {'id': 'FJ', 'value': 0.0}, {'id': 'FH', 'value': 0.0},
                {'id': 'FK', 'value': 0.0}, {'id': 'GC', 'value': 0.0}, {'id': 'FN', 'value': 0.0},
                {'id': 'FM', 'value': 0.0}, {'id': 'GB', 'value': 0.0}, {'id': 'KA', 'value': ''},
                {'id': 'KD', 'value': 0.0}, {'id': 'AH', 'value': 0.0}, {'id': 'AQ', 'value': 0.0},
                {'id': 'KB', 'value': ''}, {'id': 'KC', 'value': ''}, {'id': 'KG', 'value': ''}
            ],
            '2033F': [
                {'id': 'GX', 'value': 0}, {'id': 'GT', 'value': 0.0}, {'id': 'GV', 'value': 0.0},
                {'id': 'GY', 'value': 0}, {'id': 'GU', 'value': 0.0}, {'id': 'GW', 'value': 0.0},
                {'id': 'GA', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'GR', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'GD', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'BA', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'BG', 'occurences': [{'address': {'postal_code': '', 'city': '', 'country_code': None}, 'occurence_number': '0001'}]},
                {'id': 'BE', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'BR', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'BD', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]}
            ],
            '2033G': [
                {'id': 'GT', 'value': 1},
                {'id': 'GA', 'occurences': [{'value': '', 'occurence_number': '0001'}]}
            ],
            '2065': [
                {'id': 'PH', 'value': ''}, {'id': 'PE', 'value': ''}, {'id': 'AS', 'value': ''},
                {'id': 'AC', 'value': ''}, {'id': 'AK', 'value': ''}, {'id': 'PA', 'value': ''},
                {'id': 'PD', 'value': None}, {'id': 'AQ', 'ftx_1': ''}, {'id': 'AP', 'value': ''},
                {'id': 'HA', 'value': 0.0}, {'id': 'AT', 'value': 0.0}, {'id': 'LC', 'value': 0.0},
                {'id': 'AL', 'value': 0.0}, {'id': 'LT', 'value': 0.0}, {'id': 'LN', 'value': 0.0},
                {'id': 'AU', 'value': 0.0}, {'id': 'LW', 'value': 0.0}, {'id': 'LX', 'value': 0.0},
                {'id': 'AX', 'value': ''}, {'id': 'LQ', 'value': ''}, {'id': 'MA', 'value': ''},
                {'id': 'BY', 'value': ''}, {'id': 'LY', 'value': ''}, {'id': 'AF', 'value': ''},
                {'id': 'HD', 'value': ''}, {'id': 'AW', 'value': ''}, {'id': 'LL', 'value': ''},
                {'id': 'AM', 'value': ''}, {'id': 'BA', 'value': ''}, {'id': 'HJ', 'value': ''},
                {'id': 'BB', 'value': 0.0}, {'id': 'LV', 'value': 0.0}, {'id': 'PF', 'value': ''},
                {'id': 'BE', 'value': 0.0}, {'id': 'BF', 'value': 0.0}, {'id': 'JA', 'value': 0.0},
                {'id': 'AR', 'value': ''}, {'id': 'CA', 'value': ''}, {'id': 'BN', 'value': 'S'},
                {'id': 'EA', 'value': ''}, {'id': 'BQ', 'value': 'S'}
            ],
            '2065BIS': [
                {'id': 'AA', 'value': 0.0}, {'id': 'AB', 'value': 0.0}, {'id': 'AC', 'value': 0.0},
                {'id': 'AD', 'value': 0.0},
                {'id': 'BE', 'occurences': [{'ftx_1': '', 'occurence_number': '0001'}, {'ftx_1': '', 'occurence_number': '0002'}, {'ftx_1': '', 'occurence_number': '0003'}, {'ftx_1': '', 'occurence_number': '0004'}]},
                {'id': 'AE', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}, {'value': 0.0, 'occurence_number': '0002'}, {'value': 0.0, 'occurence_number': '0003'}, {'value': 0.0, 'occurence_number': '0004'}]},
                {'id': 'AK', 'value': 0.0}, {'id': 'AL', 'value': 0.0}, {'id': 'AJ', 'value': 0},
                {'id': 'CV', 'value': ''}, {'id': 'CX', 'value': 0.0}, {'id': 'CY', 'value': 0.0},
                {'id': 'AH', 'value': 0.0}, {'id': 'CZ', 'value': 0.0}, {'id': 'AQ', 'value': 0.0},
                {'id': 'AM', 'value': 0.0}, {'id': 'DA', 'value': 0.0}, {'id': 'AR', 'value': 0.0},
                {'id': 'AN', 'value': 0.0}, {'id': 'DB', 'value': 0.0}, {'id': 'AS', 'value': 0.0},
                {'id': 'AP', 'value': 0.0}, {'id': 'DC', 'value': 0.0}, {'id': 'AT', 'value': 0.0},
                {'id': 'AF', 'value': 0.0}, {'id': 'AG', 'value': 0.0},
                {'id': 'CA', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'CG', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'CH', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'CJ', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'CK', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'CL', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'CM', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'CN', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'CW', 'occurences': [{'value': '', 'occurence_number': '0001'}]}
            ],
            '2069RCI': [
                {'id': 'AE', 'value': ''}, {'id': 'AD', 'value': ''}, {'id': 'BE', 'value': ''},
                {'id': 'BB', 'occurences': [
                    {'value': 0.0, 'occurence_number': '0001'}, {'value': 0.0, 'occurence_number': '0002'},
                    {'value': 0.0, 'occurence_number': '0003'}, {'value': 0.0, 'occurence_number': '0004'},
                    {'value': 0.0, 'occurence_number': '0005'}, {'value': 0.0, 'occurence_number': '0006'},
                    {'value': 0.0, 'occurence_number': '0007'}, {'value': 0.0, 'occurence_number': '0008'},
                    {'value': 0.0, 'occurence_number': '0009'}, {'value': 0.0, 'occurence_number': '0010'},
                    {'value': 0.0, 'occurence_number': '0011'}, {'value': 0.0, 'occurence_number': '0012'},
                    {'value': 0.0, 'occurence_number': '0013'}, {'value': 0.0, 'occurence_number': '0014'},
                    {'value': 0.0, 'occurence_number': '0015'}, {'value': 0.0, 'occurence_number': '0016'},
                    {'value': 0.0, 'occurence_number': '0017'}, {'value': 0.0, 'occurence_number': '0018'},
                    {'value': 0.0, 'occurence_number': '0019'}
                ]},
                {'id': 'BA', 'occurences': [
                    {'value': 'VEL', 'occurence_number': '0001'}, {'value': 'AUT', 'occurence_number': '0002'},
                    {'value': 'AUT', 'occurence_number': '0003'}, {'value': 'MEC', 'occurence_number': '0004'},
                    {'value': 'CIC', 'occurence_number': '0005'}, {'value': 'COM', 'occurence_number': '0006'},
                    {'value': '2LI', 'occurence_number': '0007'}, {'value': 'AUT', 'occurence_number': '0008'},
                    {'value': 'RAC', 'occurence_number': '0009'}, {'value': 'CIN', 'occurence_number': '0010'},
                    {'value': 'AUD', 'occurence_number': '0011'}, {'value': 'CCI', 'occurence_number': '0012'},
                    {'value': 'CSV', 'occurence_number': '0013'}, {'value': 'RTD', 'occurence_number': '0014'},
                    {'value': 'REB', 'occurence_number': '0015'}, {'value': 'AUT', 'occurence_number': '0016'},
                    {'value': 'HVE', 'occurence_number': '0017'}, {'value': '3IV', 'occurence_number': '0018'},
                    {'value': 'PAM', 'occurence_number': '0019'}
                ]},
                {'id': 'AU', 'value': 0.0}, {'id': 'BP', 'value': 0.0}, {'id': 'AV', 'value': 0.0},
                {'id': 'AM', 'value': 0.0}, {'id': 'AN', 'value': 0.0}, {'id': 'AP', 'value': 0.0},
                {'id': 'BD', 'occurences': [
                    {'value': 0.0, 'occurence_number': '0001'}, {'value': 0.0, 'occurence_number': '0002'},
                    {'value': 0.0, 'occurence_number': '0003'}, {'value': 0.0, 'occurence_number': '0004'},
                    {'value': 0.0, 'occurence_number': '0005'}, {'value': 0.0, 'occurence_number': '0006'},
                    {'value': 0.0, 'occurence_number': '0007'}, {'value': 0.0, 'occurence_number': '0008'},
                    {'value': 0.0, 'occurence_number': '0009'}, {'value': 0.0, 'occurence_number': '0010'},
                    {'value': 0.0, 'occurence_number': '0011'}, {'value': 0.0, 'occurence_number': '0012'},
                    {'value': 0.0, 'occurence_number': '0013'}, {'value': 0.0, 'occurence_number': '0014'},
                    {'value': 0.0, 'occurence_number': '0015'}, {'value': 0.0, 'occurence_number': '0016'}
                ]},
                {'id': 'BC', 'occurences': [
                    {'value': 'PVM', 'occurence_number': '0001'}, {'value': 'COR', 'occurence_number': '0002'},
                    {'value': 'CIR', 'occurence_number': '0003'}, {'value': 'CRC', 'occurence_number': '0004'},
                    {'value': 'TZM', 'occurence_number': '0005'}, {'value': 'FAM', 'occurence_number': '0006'},
                    {'value': 'BIO', 'occurence_number': '0007'}, {'value': 'PHO', 'occurence_number': '0008'},
                    {'value': 'ART', 'occurence_number': '0009'}, {'value': 'CJV', 'occurence_number': '0010'},
                    {'value': 'EOM', 'occurence_number': '0011'}, {'value': 'CPE', 'occurence_number': '0012'},
                    {'value': 'PTR', 'occurence_number': '0013'}, {'value': 'RTA', 'occurence_number': '0014'},
                    {'value': 'CIO', 'occurence_number': '0015'}, {'value': 'COL', 'occurence_number': '0016'}
                ]},
                {'id': 'BK', 'value': 0.0}, {'id': 'AT', 'ftx_1': ''},
                {'id': 'AW', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'AX', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'AY', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'AZ', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'AH', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]},
                {'id': 'AL', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'AS', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'BL', 'occurences': [{'value': '', 'occurence_number': '0001'}]},
                {'id': 'BN', 'occurences': [{'value': 0.0, 'occurence_number': '0001'}]}
            ]
        }
        for report_name, config in REPORTS_TO_EXPORT.items():
            lines = []
            for report in config['reports']:
                section_id = self.env.ref(f'l10n_fr_reports.{report}').id
                section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
                section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
                self._add_section(section, section_options, config)
                lines.extend(section._get_lines(section_options))

            mapper = self.env[config['handler']]._get_code_to_edi_id()
            vals = self.env['l10n_fr_reports.fiscal_declaration_handler']._export_report_data(lines, mapper)

            with self.subTest(vals=vals):
                self.assertEqual(vals, expected[report_name])

    ###############################################################
    # Fiscale reports add/remove lines
    ###############################################################

    def test_simple_add_lines(self):
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        config = REPORTS_TO_EXPORT.get('2033G')
        section_id = self.env.ref('l10n_fr_reports.l10n_fr_2033_G').id
        section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
        section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
        self._add_section(section, section_options, config)

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                    Balance
            [  0,                                                        1],
            [
                ('Subsidiaries and Investments',                        ''),
                ('Total number of subsidiaries held by the company',     1),
                ('Add subsidiary',                                      ''),
                ('Subsidiary 1',                                        ''),
            ],
            section_options
        )

    def test_simple_remove_lines(self):
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        config = REPORTS_TO_EXPORT.get('2033G')
        section_id = self.env.ref('l10n_fr_reports.l10n_fr_2033_G').id
        section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
        section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
        pre = [line.code for line in section._get_lines(section_options)]
        self._add_section(section, section_options, config)
        lines = [line for line in section._get_lines(section_options) if line.code not in pre]

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                    Balance
            [  0,                                                        1],
            [
                ('Subsidiaries and Investments',                        ''),
                ('Total number of subsidiaries held by the company',     1),
                ('Add subsidiary',                                      ''),
                ('Subsidiary 1',                                        ''),
            ],
            section_options
        )

        self._remove_section(lines, section, section_options, config)

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                    Balance
            [  0,                                                        1],
            [
                ('Subsidiaries and Investments',                        ''),
                ('Total number of subsidiaries held by the company',     0),
                ('Add subsidiary',                                      ''),
            ],
            section_options
        )

    def test_multi_add_lines(self):
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        config = REPORTS_TO_EXPORT.get('2033G')
        section_id = self.env.ref('l10n_fr_reports.l10n_fr_2033_G').id
        section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
        section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
        for i in range(3):
            self._add_section(section, section_options, config)

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                    Balance
            [  0,                                                        1],
            [
                ('Subsidiaries and Investments',                        ''),
                ('Total number of subsidiaries held by the company',     3),
                ('Add subsidiary',                                      ''),
                ('Subsidiary 1',                                        ''),
                ('Subsidiary 2',                                        ''),
                ('Subsidiary 3',                                        ''),
            ],
            section_options
        )

    def test_multi_remove_lines(self):
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        config = REPORTS_TO_EXPORT.get('2033G')
        section_id = self.env.ref('l10n_fr_reports.l10n_fr_2033_G').id
        section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
        section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
        pre = [line.code for line in section._get_lines(section_options)]
        for i in range(3):
            self._add_section(section, section_options, config)
        lines = [line for line in section._get_lines(section_options) if line.code not in pre]

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                    Balance
            [  0,                                                        1],
            [
                ('Subsidiaries and Investments',                        ''),
                ('Total number of subsidiaries held by the company',     3),
                ('Add subsidiary',                                      ''),
                ('Subsidiary 1',                                        ''),
                ('Subsidiary 2',                                        ''),
                ('Subsidiary 3',                                        ''),
            ],
            section_options
        )

        self._remove_section(lines, section, section_options, config, 'Subsidiary 2')  # Subsidiary 3 will be renamed to Subsidiary 2
        lines = [line for line in section._get_lines(section_options) if line.code not in pre]
        self._remove_section(lines, section, section_options, config, 'Subsidiary 2')

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                    Balance
            [  0,                                                        1],
            [
                ('Subsidiaries and Investments',                        ''),
                ('Total number of subsidiaries held by the company',     1),
                ('Add subsidiary',                                      ''),
                ('Subsidiary 1',                                        ''),
            ],
            section_options
        )

    def test_middle_remove_multi_lines(self):
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        config = REPORTS_TO_EXPORT.get('2033G')
        section_id = self.env.ref('l10n_fr_reports.l10n_fr_2033_G').id
        section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
        section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
        pre = [line.code for line in section._get_lines(section_options)]
        for i in range(3):
            self._add_section(section, section_options, config)
        lines = [line for line in section._get_lines(section_options) if line.code not in pre]

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                    Balance
            [  0,                                                        1],
            [
                ('Subsidiaries and Investments',                        ''),
                ('Total number of subsidiaries held by the company',     3),
                ('Add subsidiary',                                      ''),
                ('Subsidiary 1',                                        ''),
                ('Subsidiary 2',                                        ''),
                ('Subsidiary 3',                                        ''),
            ],
            section_options
        )

        self._remove_section(lines, section, section_options, config, 'Subsidiary 2')

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                    Balance
            [  0,                                                        1],
            [
                ('Subsidiaries and Investments',                        ''),
                ('Total number of subsidiaries held by the company',     2),
                ('Add subsidiary',                                      ''),
                ('Subsidiary 1',                                        ''),
                ('Subsidiary 2',                                        ''),
            ],
            section_options
        )

    def test_simple_add_lines_aggregation_report(self):
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        config = REPORTS_TO_EXPORT.get('2033C')
        section_id = self.env.ref('l10n_fr_reports.l10n_fr_2033_C_2').id
        section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
        section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
        self._add_section(section, section_options, config)

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                                             Asset value      Depriciation       Residual value
            [  0,                                                                                   2,               3,                  4],
            [
                ('Fixed Assets – Depreciation – Capital Gains – Capital Losses, part. 2',           '',              '',                ''),
                ('III – Capital Gains & Capital Losses',                                            '',              '',                ''),
                ('Fixed assets',                                                                    '',              '',                ''),
                ('Add asset',                                                                       '',              '',                ''),
                ('TOTAL',                                                                            0,               0,                 0),
                ('Capital gains taxable at 19%',                                                    '',              '',                 0),
                ('Adjustments',                                                                     '',              '',                ''),
                ('TOTAL',                                                                           '',              '',                ''),
                ('Asset 1',                                                                          0,               0,                 0),
            ],
            section_options
        )

    def test_simple_remove_lines_aggregation_report(self):
        options = self._generate_fiscal_options()
        self.env['l10n_fr_reports.send.fiscal.declaration'].create({
            'report_id': self.env.ref('l10n_fr_reports.fiscal_declaration').id,
            'date_from': REPORT_PERIOD_START_DATE,
            'date_to': REPORT_PERIOD_END_DATE,
        })

        fiscal_declaration = self.env['account.report'].browse(options['sections_source_id'])
        config = REPORTS_TO_EXPORT.get('2033C')
        section_id = self.env.ref('l10n_fr_reports.l10n_fr_2033_C_2').id
        section = fiscal_declaration.section_report_ids.filtered(lambda sect: sect.id == section_id)
        section_options = section.get_options({'date': options.get('date'), 'unfold_all': True})
        pre = [line.code for line in section._get_lines(section_options)]
        self._add_section(section, section_options, config)
        lines = [line for line in section._get_lines(section_options) if line.code not in pre]

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                                             Asset value      Depriciation       Residual value
            [  0,                                                                                   2,               3,                  4],
            [
                ('Fixed Assets – Depreciation – Capital Gains – Capital Losses, part. 2',           '',              '',                ''),
                ('III – Capital Gains & Capital Losses',                                            '',              '',                ''),
                ('Fixed assets',                                                                    '',              '',                ''),
                ('Add asset',                                                                       '',              '',                ''),
                ('TOTAL',                                                                            0,               0,                 0),
                ('Capital gains taxable at 19%',                                                    '',              '',                 0),
                ('Adjustments',                                                                     '',              '',                ''),
                ('TOTAL',                                                                           '',              '',                ''),
                ('Asset 1',                                                                          0,               0,                 0),
            ],
            section_options
        )

        self._remove_section(lines, section, section_options, config)

        self.assertLinesValues(
            section._get_lines(section_options),
            # Name                                                                             Asset value      Depriciation       Residual value
            [  0,                                                                                   2,               3,                  4],
            [
                ('Fixed Assets – Depreciation – Capital Gains – Capital Losses, part. 2',           '',              '',                ''),
                ('III – Capital Gains & Capital Losses',                                            '',              '',                ''),
                ('Fixed assets',                                                                    '',              '',                ''),
                ('Add asset',                                                                       '',              '',                ''),
                ('TOTAL',                                                                            0,               0,                 0),
                ('Capital gains taxable at 19%',                                                    '',              '',                 0),
                ('Adjustments',                                                                     '',              '',                ''),
                ('TOTAL',                                                                           '',              '',                ''),
            ],
            section_options
        )

    def test_export_template_fiscal_declaration(self):
        import_wizard = self.env['l10n_fr.import.fiscal.declaration'].create({
            'report_id': self.report.id,
        })
        action = import_wizard.action_download_template()

        # Just to see if the function returned something
        self.assertEqual(action['type'], 'ir.actions.act_url')

        attachment = self.env['ir.attachment'].search([
            ('name', '=', 'template_fiscal_declaration.xlsx')
        ], limit=1)
        self.assertTrue(attachment)
        workbook = load_workbook(
            filename=io.BytesIO(attachment.raw),
            data_only=True,
        )
        # Just to make sure we have different sheets
        self.assertTrue(workbook.sheetnames)

        sheet = workbook[workbook.sheetnames[0]]
        headers = [cell.value for cell in sheet[1]]

        self.assertEqual(headers, [
            'Line code',
            'Line name',
            'Expression label',
            'Expected format',
            'External Value',
            'Comment',
        ])

    def test_import_template_fiscal_declaration(self):
        import_file = file_open('l10n_fr_reports/tests/expected_files/template_fiscal_declaration_2033_B.xlsx', 'rb').read()
        report_2033_B = self.env.ref('l10n_fr_reports.l10n_fr_2033_B')
        # Needed because we cannot edit external value with multiple companies
        self.env = self.env(context=dict(self.env.context, allowed_company_ids=self.env.company.ids))

        report_options = self._generate_options(report_2033_B, '2026-01-01', '2026-12-31')
        report_line = [line for line in report_2033_B._get_lines(report_options) if line.code == 'FR_2033_B_eq_lease']
        self.assertEqual(report_line[0].columns[0].no_format, 0.0)

        wizard = self.env['l10n_fr.import.fiscal.declaration'].create({
            'data': BinaryBytes(import_file),
            'filename': 'import_file.xlsx',
            'report_id': self.report.id,
        })

        action = wizard.action_import_template()
        # Just to see if the function returned something
        self.assertEqual(action['type'], 'ir.actions.client')
        self.env.flush_all()

        # Check if the external value has changed
        report_line = [line for line in report_2033_B._get_lines(report_options) if line.code == 'FR_2033_B_eq_lease']
        self.assertEqual(report_line[0].columns[0].no_format, 400.0)
