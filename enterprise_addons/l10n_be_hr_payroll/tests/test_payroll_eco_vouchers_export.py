# Part of Odoo. See LICENSE file for full copyright and licensing details.

import io
from datetime import date
from odoo.tests import tagged, HttpCase
from odoo.exceptions import UserError
from .common import TestPayrollCommon

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


@tagged('post_install_l10n', 'post_install', '-at_install')
class TestPayrollEcoVouchersExport(TestPayrollCommon, HttpCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.admin_user = cls.env.ref('base.user_admin')
        cls.admin_user.company_ids |= cls.belgian_company
        cls.admin_user.company_id = cls.belgian_company
        cls.cp200 = cls.env.ref('l10n_be_hr_payroll.l10n_be_joint_committee_200')
        cls.test_contracts.write({'l10n_be_joint_committee_id': cls.cp200.id})

    def test_eco_vouchers_export(self):
        self.authenticate('admin', 'admin')

        payslip = self.env['hr.payslip'].create({
            'employee_id': self.employee_test.id,
            'version_id': self.employee_test.version_id.id,
            'date_from': date(2026, 6, 1),
            'date_to': date(2026, 6, 30),
            'company_id': self.belgian_company.id,
        })
        payslip.compute_sheet()
        payslip.action_payslip_done()
        wizard = self.env['l10n.be.eco.vouchers.wizard'].with_company(self.belgian_company).create({
            'company_id': self.belgian_company.id,
            'reference_year': 2026,
        })

        self.assertTrue(wizard.line_ids, "Wizard should have lines for the employees created in common")

        response = self.url_open(f'/export/ecovouchers/{wizard.id}')
        self.assertEqual(response.status_code, 200)

        if load_workbook:
            output = io.BytesIO(response.content)
            workbook = load_workbook(output)
            sheet = workbook['Worksheet']

            expected_headers = [
                "National registration number (e.g. 790227 183 12)",
                "Employee last name (e.g. jansen)",
                "Employee first name (e.g. max)",
                "Your internal employee number (e.g. 152d97)",
                "Number voucher [a] (e.g. 18)",
                "Value voucher [b] (e.g. 5.5)",
                "Total [a] x [b] (e.g. 99)",
                "Employee birth date (dd/mm/yyyy)",
                "Employee gender (m/f)",
                "Employee language (nl/fr/en)",
                "Cost center (e.g. cc1. maximum 10 characters)",
                "Your company number  (e.g. be 0834013324)",
                "Delivery address street (e.g. av. des volontaires)",
                "Delivery address number (e.g. 19)",
                "Delivery address box  (e.g. b2)",
                "Delivery address zipcode (e.g. 1160)",
                "Delivery address city (e.g. oudergem)",
                "Contract status"
            ]

            actual_headers = [cell.value for cell in sheet[1]]
            self.assertEqual(actual_headers, expected_headers)

            self.assertTrue(sheet.max_row > 1, "There should be at least one data row")

            # Check the 'Statut contrat' column
            last_col_index = len(expected_headers)
            status_value = sheet.cell(row=2, column=last_col_index).value
            self.assertIn(status_value, ['Active', 'End of collaboration'])

    def test_report_can_only_be_generated_from_root_company(self):
        companies = self.multibranch_company
        with self.assertRaises(UserError):
            self.env['l10n.be.eco.vouchers.wizard'].with_company(companies[1]).create({})

    def test_include_child_branches_in_report(self):
        companies = self.belgian_company | self.multibranch_company
        # company 0 -> root / no child
        # company 1 (is parent of) company 2 (is parent of) company 3
        target_company = companies[1]

        employees = self.create_employee([
            {
                "name": f'Employee {company.name}',
                "company_id": company.id,
                'contract_date_start': date(2026, 1, 1),
                'l10n_be_joint_committee_id': self.env.ref('l10n_be_hr_payroll.l10n_be_joint_committee_200').id,
                'l10n_be_salary_scale_id': self.env.ref('l10n_be_hr_payroll.cp200_b').id,
            } for company in companies
        ])
        self.create_and_validate_payslips(employees=employees, year=2026, months=[6])

        wizard = self.env['l10n.be.eco.vouchers.wizard'].with_company(target_company).create({
            'company_id': target_company.id,
            'reference_year': 2026,
        })

        expected_companies = companies[1:]
        self.assertEqual(wizard.branch_ids, expected_companies,
                         "Report should include all three levels of the hierarchy.")

        expected_reported_employees = employees[1:]
        reported_employees = wizard.line_ids.mapped('employee_id')
        self.assertEqual(reported_employees, expected_reported_employees,
                         "Report should include all three employees")

        self.assertEqual(len(wizard.line_ids), 3,
                         "There should be exactly one line per employee across all branches.")
