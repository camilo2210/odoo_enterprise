import unittest

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from datetime import date

from odoo.addons.l10n_in_hr_payroll.tests.common import TestPayrollCommon
from odoo.tests import tagged


@tagged('post_install_l10n', 'post_install', '-at_install')
class TestHrEPFReport(TestPayrollCommon):

    def test_hr_epf_report(self):
        """ Test EPF report XLSX generation and values."""
        self.jethalal_emp.write({
            'wage': 10000,
            'l10n_in_basic_salary_amount': 1000,
            'l10n_in_uan': 5674839203,
        })
        payslips = self.env['hr.payslip'].create([{
                'name': 'Jethalal Payslip',
                'employee_id': self.jethalal_emp.id,
                'version_id': self.contract_jethalal.id,
                'date_from': date(2023, 1, 1),
                'date_to': date(2023, 1, 31),
            }, {
                'name': 'Rahul Payslip',
                'employee_id': self.rahul_emp.id,
                'version_id': self.contract_rahul.id,
                'date_from': date(2023, 1, 1),
                'date_to': date(2023, 1, 31),
            }
        ])

        payslips.compute_sheet()
        payslips.action_payslip_done()

        epf_report = self.env['l10n.in.hr.payroll.epf.report'].create({
            'month': '1',
            'year': '2023',
            'export_report_type': 'report',
        })

        # Action to generate the XLSX file
        epf_report.action_export_xlsx()
        self.assertTrue(epf_report.xls_file, "The XLS file was not generated.")

        if load_workbook is None:
            raise unittest.SkipTest("openpyxl not available")

        with epf_report.xls_file.open() as data:
            xlsx = load_workbook(data)
        sheet = xlsx.worksheets[0]
        sheet_values = list(sheet.values)

        expected_values = {
            0: ['EPF Statement January-2023', None, None, None, None, None, None, None, None, None, None, None],
            1: ['Name of Establishment', 'Company IN', None, None, None, None, None, None, None, None, None, None],
            2: ['Establishment ID', None, None, None, None, None, None, None, None, None, None, None],
            3: ['Total Members', '2', None, None, None, None, None, None, None, None, None, None],
            # Header
            4: [
                'SI NO',
                'UAN',
                'MEMBER NAME',
                'GROSS WAGES',
                'EPF WAGES',
                'EPS WAGES',
                'EDLI WAGES',
                'EPF CONTRIBUTION REMITTED',
                'EPS CONTRIBUTION REMITTED',
                'EPF EPS DIFFERENCE REMITTED',
                'NCP DAYS',
                'REFUNDED OF ADVANCES'
            ],
            # Employee Rows
            5: ['1', '5674839203', 'Jethalal', 10000, 1000, 1000, 1000, 120, 83.3, 36.7, 0, 0],
            6: ['2', None, 'Rahul', 5000, 1750, 1750, 1750, 210, 145.78, 64.22, 0, 0],
        }
        for row, values in expected_values.items():
            for row_value, expected_value in zip(sheet_values[row], values):
                self.assertEqual(row_value, expected_value)

    def test_hr_epf_summary(self):
        """ Test EPF summary XLSX generation and values."""
        self.jethalal_emp.write({
            'wage': 10000,
            'l10n_in_basic_salary_amount': 1000,
            'registration_number': 9876543210,
            'l10n_in_pf_account_number': 1234567890,
            'l10n_in_uan': 5674839203,
        })
        self.rahul_emp.write({
            'registration_number': False,
            'l10n_in_pf_account_number': 1029384756,
            'l10n_in_uan': 564839479,
        })
        payslips = self.env['hr.payslip'].create([{
                'name': 'Jethalal Payslip',
                'employee_id': self.jethalal_emp.id,
                'version_id': self.contract_jethalal.id,
                'date_from': date(2023, 1, 1),
                'date_to': date(2023, 1, 31),
            }, {
                'name': 'Rahul Payslip',
                'employee_id': self.rahul_emp.id,
                'version_id': self.contract_rahul.id,
                'date_from': date(2023, 1, 1),
                'date_to': date(2023, 1, 31),
            }
        ])

        payslips.compute_sheet()
        payslips.action_payslip_done()

        epf_summary = self.env['l10n.in.hr.payroll.epf.report'].create({
            'month': '1',
            'year': '2023',
            'export_report_type': 'summary',
        })

        # Action to generate the XLSX file
        epf_summary.action_export_xlsx()
        self.assertTrue(epf_summary.xls_file, "The XLS file was not generated.")

        if load_workbook is None:
            raise unittest.SkipTest("openpyxl not available")

        with epf_summary.xls_file.open() as data:
            xlsx = load_workbook(data)
        sheet = xlsx.worksheets[0]
        sheet_values = list(sheet.values)

        # Expected rows (data starts from row index 6)
        expected_values = {
            6: ['9876543210', 'Jethalal', '1234567890', '5674839203', 1000, 120, 0, 120, 83.3, 0, 5, 328.3],
            7: [None, 'Rahul', '1029384756', '564839479', 1750, 210, 0, 210, 145.78, 0, 8.75, 574.53],
        }
        for row, values in expected_values.items():
            for row_value, expected_value in zip(sheet_values[row], values):
                self.assertEqual(row_value, expected_value)
