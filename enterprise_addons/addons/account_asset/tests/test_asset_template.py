import io
import unittest
from datetime import datetime

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from odoo.tests import HttpCase, tagged
from odoo.tools import BinaryBytes


@unittest.skipIf(load_workbook is None, "openpyxl not available")
@tagged('at_install', '-post_install')  # LEGACY at_install
class TestAssetTemplate(HttpCase):

    def test_download_asset_template(self):
        def get_downloaded_template_rows_with_output(company):
            response = self.url_open(f'/web/binary/download_asset_template/{company.id}')
            self.assertEqual(response.status_code, 200)

            output = io.BytesIO(response.content)
            workbook = load_workbook(output)
            sheet = workbook['Assets']
            actual_rows = [
                tuple('' if v is None else str(v) for v in actual_row[:17])
                for actual_row in sheet.iter_rows(min_row=2, values_only=True)
            ]
            return actual_rows, output

        def get_expected_rows(fixed_asset_account_name):
            return [
                ('Computer 1', '800.00', '01-01-2025', '', depreciation_model.display_name, 'No Prorata', '01-01-2025', '200.00', '0.00', company.name, fixed_asset_account_name, '', ''),
                ('Computer 2', '25000.00', '03-15-2025', '', depreciation_model.display_name, 'Based on days per period', '03-15-2025', '0.00', '2000.00', company.name, fixed_asset_account_name, '', ''),
                ('Machine A', '15000.00', '09-01-2024', '', depreciation_model.display_name, 'Constant Periods', '09-01-2024', '1500.00', '500.00', company.name, fixed_asset_account_name, '', ''),
                ('Machine B', '100000.00', '06-20-2025', '', depreciation_model.display_name, 'Constant Periods', '06-20-2025', '10000.00', '10000.00', company.name, fixed_asset_account_name, '', ''),
            ]

        self.authenticate('admin', 'admin')
        company = self.env.company
        account_model = self.env['account.account'].with_company(company)

        # Archive all existing accounts to control the test environment
        account_model.search([]).write({'active': False})

        journal_id = self.env['account.journal'].create({
            'name': 'Test Miscellaneous Operations',
            'code': 'MISC-t',
            'type': 'general',
        })

        depreciation_account = account_model.create({
            'name': 'Depreciation Account',
            'code': 'DP001',
            'account_type': 'asset_fixed',
        })

        expense_account = account_model.create({
            'name': 'Depreciation Account',
            'code': 'EX001',
            'account_type': 'expense',
        })

        depreciation_model = self.env['account.depreciation.model'].create({
            'method': 'linear',
            'method_number': 5,
            'method_period': '12',
            'journal_id': journal_id.id,
            'company_id': company.id,
        })
        conflicting_models = self.env['account.depreciation.model'].search([
            ('display_name', '=', depreciation_model.display_name),
            ('id', '!=', depreciation_model.id)
        ])
        conflicting_models.write({'active': False})

        asset_account = account_model.create({
            'name': 'Fixed Asset Account',
            'code': 'FA001',
            'account_type': 'asset_fixed',
            'depreciation_model_id': depreciation_model.id,
            'asset_depreciation_account_id': depreciation_account.id,
            'asset_expense_account_id': expense_account.id,
        })

        actual_rows, output = get_downloaded_template_rows_with_output(company)

        expected_rows = get_expected_rows(fixed_asset_account_name='FA001 Fixed Asset Account')
        self.assertEqual(actual_rows, expected_rows)

        import_wizard = self.env['base_import.import'].create({
            'res_model': 'account.asset',
            'file': BinaryBytes(output.getvalue()),
            'file_type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        result = import_wizard.parse_preview({
            'has_headers': True,
        })
        # Check that there were no error messages in the result
        self.assertFalse(result.get('messages'), "Import failed with messages: %s" % result.get('messages'))

        options = {
            'separator': ',',
            'quoting': '"',
            'date_format': '%m-%d-%Y',
            'headers': True,
            'has_headers': True,
        }
        asset_fields = ['name', 'original_value', 'acquisition_date', 'asset_group_id', 'model_id',
                        'prorata_computation_type', 'prorata_date', 'already_depreciated_amount_import', 'salvage_value', 'company_id',
                        'account_asset_id', 'import_account_depreciation_id', 'import_account_depreciation_expense_id']
        result = import_wizard.execute_import(
            fields=asset_fields,
            columns=[],
            options=options
        )

        # Check that there were no error messages in the result
        self.assertFalse(result.get('messages'), "Import failed with messages: %s" % result.get('messages'))

        # Fetch the created asset
        assets_imported = self.env['account.asset'].search([('model_id', '=', depreciation_model.id)]).sorted('name')
        self.assertEqual(len(assets_imported), 4)

        # Check asset values
        expected_vals = []
        for expected_row in expected_rows:
            expected_vals.append({
                'name': expected_row[0],
                'original_value': float(expected_row[1]),
                'acquisition_date': datetime.strptime(expected_row[2], '%m-%d-%Y').date(),
                'model_id': depreciation_model.id,
                'prorata_date': datetime.strptime(expected_row[6], '%m-%d-%Y').date(),
                'already_depreciated_amount_import': float(expected_row[7]),
                'salvage_value': float(expected_row[8]),
                'company_id': company.id,
                'account_asset_id': asset_account.id,
                'journal_id': journal_id.id,
            })
        self.assertRecordValues(assets_imported, expected_vals)

        # Check asset template download when the fixed asset account has no code.
        asset_account.code = False
        actual_rows, _ = get_downloaded_template_rows_with_output(company)

        expected_rows = get_expected_rows(fixed_asset_account_name='Fixed Asset Account')
        self.assertEqual(actual_rows, expected_rows)
